"""싼타페 페이스리프트(MX5 F/L) 가격 언급 게시물 수집 - 엑셀 양식(price-template.xlsx) 만들기.

서체·크기·테두리는 기존 보고 양식(seat-extension/seat-template.xlsx)을 그대로 따른다:
현대산스 Text 11pt · 3행 머리글(가운데) · 얇은 테두리 · A4 틀 고정 · 머리글 필터. 머리글 행은 현대 블루 바탕에 흰색 굵은 글씨.
양식 파일을 바탕으로 열어 그 셀 서식을 복사해 쓰므로, 기존 양식과 같은 글꼴·테마가 들어간다.

시트 3개 (각 시트의 서식 견본 행은 받을 때 price.js 가 내용으로 바꾼다)
1. 게시물: 3행 머리글 · 4행 서식 견본
   B 순번 · C 동호회명 · D 작성일 · E 제목 · F 내용 요약 · G URL · H 작성자 닉네임 · I 조회수 · J 댓글수
   · K 구간 · L 예상 가격 여론 · M 예상 가격·인상 폭 언급 · N 예상 인상 폭(만 원) · O 비교 차종 · P 구매 영향 · Q 여론 요지
2. 구간별 여론: 1행 제목 · 2행 조건 · 3행 여론 변화 보고 · 5행 머리글 · 6행 서식 견본
3. 주별 추이: 1행 제목 · 2행 조건 · 3행 머리글 · 4행 서식 견본

    python scripts/build_price_template.py
"""
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
BASE = ROOT / "seat-extension" / "seat-template.xlsx"
OUT = ROOT / "price-extension" / "price-template.xlsx"
HYUNDAI_BLUE = "FF002C5F"  # 머리글 행: 현대 블루 바탕 · 흰색 굵은 글씨
RED = "FFE00000"  # 기존 양식의 N 빨강
BLUE = "FF0070C0"  # 기존 양식의 Y 파랑
GREEN = "FF00B050"  # 중립 녹색

# (머리글, 너비, 서식, 모양) 모양: c 가운데 · w 왼쪽 줄바꿈 · l 왼쪽 · u 링크
POSTS = [
    ("순번", 5.5, "General", "c"),
    ("동호회명", 16, "General", "c"),
    ("작성일", 17.7, "yyyy-mm-dd hh:mm", "c"),
    ("제목", 40, "General", "w"),
    ("내용 요약", 60, "General", "w"),
    ("URL", 38, "General", "u"),
    ("작성자 닉네임", 16, "General", "c"),
    ("조회수", 9, "#,##0", "c"),
    ("댓글수", 9, "#,##0", "c"),
    ("구간", 20, "General", "c"),
    ("예상 가격 여론", 11, "General", "c"),
    ("예상 가격·인상 폭 언급", 34, "General", "w"),
    ("예상 인상 폭(만 원)", 11, "#,##0", "c"),
    ("비교 차종", 13, "General", "c"),
    ("구매 영향", 16, "General", "c"),
    ("여론 요지", 50, "General", "w"),
]
STAGES = [
    ("구간", 22, "General", "c"),
    ("기간", 25, "General", "c"),
    ("일수", 7, "#,##0", "c"),
    ("글 수", 8, "#,##0", "c"),
    ("하루 평균 글 수", 10, "0.0", "c"),
    ("긍정", 7, "#,##0", "c"),
    ("중립", 7, "#,##0", "c"),
    ("부정", 7, "#,##0", "c"),
    ("긍정 비율", 10, "0%", "c"),
    ("부정 비율", 10, "0%", "c"),
    ("예상 가격 언급 글 수", 11, "#,##0", "c"),
    ("예상 인상 폭 평균(만 원)", 12, "#,##0", "c"),
    ("예상 인상 폭 중앙값(만 원)", 12, "#,##0", "c"),
    ("아반떼 비교 언급", 11, "#,##0", "c"),
    ("투싼 비교 언급", 11, "#,##0", "c"),
    ("구매 대기·관망", 11, "#,##0", "c"),
    ("타 차종 고려·이탈", 11, "#,##0", "c"),
    ("평균 조회수", 11, "#,##0", "c"),
    ("평균 댓글수", 11, "#,##0.0", "c"),
    ("여론 요지", 70, "General", "w"),
]
WEEKS = [
    ("주", 25, "General", "c"),
    ("구간", 22, "General", "c"),
    ("이 주의 사건", 22, "General", "c"),
    ("글 수", 8, "#,##0", "c"),
    ("긍정", 7, "#,##0", "c"),
    ("중립", 7, "#,##0", "c"),
    ("부정", 7, "#,##0", "c"),
    ("부정 비율", 10, "0%", "c"),
    ("예상 인상 폭 평균(만 원)", 12, "#,##0", "c"),
    ("아반떼 비교 언급", 11, "#,##0", "c"),
    ("투싼 비교 언급", 11, "#,##0", "c"),
    ("평균 조회수", 11, "#,##0", "c"),
    ("평균 댓글수", 11, "#,##0.0", "c"),
]


def styler(base):
    """기존 양식 셀의 서식(글꼴·테두리·맞춤)을 그대로 가져온다."""
    head = base["B3"]
    center = base["B4"]  # 가운데
    wrap = base["F4"]  # 왼쪽 · 줄바꿈
    link = base["G4"]  # 링크 글자색

    def put(cell, src, fmt=None, bold=None):
        cell.font = copy(src.font)
        if bold is not None:
            f = copy(src.font)
            f.b = bold
            cell.font = f
        cell.border = copy(src.border)
        cell.alignment = copy(src.alignment)
        cell.fill = copy(src.fill)
        cell.number_format = fmt or src.number_format

    def header(cell, name):
        cell.value = name
        put(cell, head)
        f = copy(head.font)
        f.b = True
        f.color = "FFFFFFFF"
        cell.font = f
        cell.fill = PatternFill("solid", fgColor=HYUNDAI_BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def body(cell, kind, fmt):
        put(cell, {"c": center, "w": wrap, "l": wrap, "u": link}[kind], fmt, bold=False)
        if kind == "l":
            cell.alignment = Alignment(vertical="center")

    return header, body, put, head


def table(ws, columns, head_row, header, body):
    for i, (name, width, fmt, kind) in enumerate(columns):
        col = get_column_letter(2 + i)
        ws.column_dimensions[col].width = width
        header(ws[f"{col}{head_row}"], name)
        body(ws[f"{col}{head_row + 1}"], kind, fmt)
    last = get_column_letter(1 + len(columns))
    ws.freeze_panes = f"A{head_row + 1}"
    ws.auto_filter.ref = f"B{head_row}:{last}{head_row + 1}"
    return last


def tone_colors(ws, ref, first):
    """예상 가격 여론 글자색: 부정 빨강 · 긍정 파랑 (기존 양식의 Y 파랑 · N 빨강과 같은 색) · 중립 녹색"""
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f'{first}="부정"'], font=Font(color=RED, bold=True)))
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f'{first}="긍정"'], font=Font(color=BLUE, bold=True)))
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f'{first}="중립"'], font=Font(color=GREEN, bold=True)))


def build():
    wb = load_workbook(BASE)
    ws = wb.worksheets[0]
    ws.title = "게시물"
    # 기존 양식의 셀 서식 견본을 따로 둔 뒤 시트를 비운다
    samples = {k: copy(ws[k]._style) for k in ["B3", "B4", "F4", "G4"]}
    ref = wb.create_sheet("_견본")
    for k, st in samples.items():
        ref[k]._style = st
    ref["G4"].value = "https://cafe.naver.com/"
    header, body, put, head = styler(ref)
    for row in ws.iter_rows():
        for c in row:
            c.value = None
            c.hyperlink = None
            c.style = "Normal"
    ws.conditional_formatting = ConditionalFormattingList()
    for col in list(ws.column_dimensions):
        del ws.column_dimensions[col]

    table(ws, POSTS, 3, header, body)
    # URL 견본 셀에 링크를 걸어 둬야 양식에 시트 관계 파일이 생긴다(받을 때 글 주소마다 링크를 거는 데 필요)
    ws["G4"] = "https://cafe.naver.com/"
    ws["G4"].hyperlink = "https://cafe.naver.com/"
    tone_colors(ws, "L4:L20000", "L4")

    def titled(name, title, columns, head_row):
        s = wb.create_sheet(name)
        s["B1"] = title
        put(s["B1"], head)
        s["B1"].border = copy(ref["A1"].border)
        s["B1"].alignment = Alignment(vertical="center")
        s["B2"] = "{{conditions}}"
        body(s["B2"], "l", "General")
        s["B2"].border = copy(ref["A1"].border)
        last = table(s, columns, head_row, header, body)
        return s, last

    s, last = titled("구간별 여론", "싼타페 F/L 예상 가격 여론 변화 (기준 사건 전후)", STAGES, 5)
    s.merge_cells(f"B3:{last}3")
    s["B3"] = "{{overall}}"
    body(s["B3"], "w", "General")
    s["B3"].alignment = Alignment(vertical="top", wrap_text=True)
    s.row_dimensions[3].height = 150

    titled("주별 추이", "싼타페 F/L 예상 가격 여론 주별 추이", WEEKS, 3)

    # 기존 양식의 보안 문구 바닥글을 세 시트 모두에
    for sheet in wb.worksheets:
        if sheet is not ws and sheet is not ref:
            sheet.oddFooter.left.text = ws.oddFooter.left.text
            sheet.oddFooter.left.font = ws.oddFooter.left.font
            sheet.oddFooter.left.size = ws.oddFooter.left.size
            sheet.oddFooter.left.color = ws.oddFooter.left.color

    wb.remove(ref)
    wb.save(OUT)
    print(OUT.relative_to(ROOT))


if __name__ == "__main__":
    build()
