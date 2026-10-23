"""시트 위험 게시물 수집 - 보고 양식(seat-template.xlsx)의 열 너비·행 높이 맞추기.

모든 글 행은 같은 높이(ROW_HEIGHT)로 나오게 한다. 프로그램은 4행(양식 행)의 높이·서식을 글 행마다 그대로 복사한다.
- 제목·URL 은 가능한 한 한 줄에 들어가도록 넓게 (제목이 아주 길면 칸 안에서 줄바꿈)
- 주요내용 요약(최대 120자)·누락 사유(앞에 붙는 로데이터 작성 전 문구 포함)는 4줄 안에 들어가는 너비
- 모든 칸 세로 가운데 정렬

이 스크립트를 고친 뒤에는 최종 파일 양식도 다시 만든다:
    python scripts/build_seat_template.py && python scripts/build_seat_final_template.py
"""
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment

ROOT = Path(__file__).resolve().parent.parent / "seat-extension"
ROW_HEIGHT = 66  # 4줄 (11pt)
HEAD_HEIGHT = 36
# B 순번 · C 차종 · D 동호회명 · E 게시일 · F 제목 · G URL · H 주요내용 요약 · I 포함 여부 · J 누락 사유 · K 주제 관련도
WIDTHS = {"A": 2, "B": 6, "C": 7, "D": 22, "E": 12, "F": 56, "G": 46, "H": 68, "I": 14, "J": 80, "K": 11}
CENTER = "BCDEIK"
WRAP = "FHJ"


def build():
    path = ROOT / "seat-template.xlsx"
    wb = load_workbook(path)
    ws = wb.active
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[3].height = HEAD_HEIGHT
    ws.row_dimensions[4].height = ROW_HEIGHT
    for col in "BCDEFGHIJK":
        h = ws[f"{col}3"]
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c = ws[f"{col}4"]
        c.alignment = Alignment(vertical="center", wrap_text=col in WRAP, horizontal="center" if col in CENTER else None)
    wb.save(path)
    print("saved", path)


if __name__ == "__main__":
    build()
