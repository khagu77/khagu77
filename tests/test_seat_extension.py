"""'LX3 / ME1 / MX5 전동 시트 작동 위험 상황 관련 게시물 수집' 확장프로그램(seat-extension) 통합 테스트.

cafe.naver.com 을 로컬 HTTPS 서버로 돌려 카페 첫 화면·검색 결과·게시글을 흉내 내고,
게시물 찾기 → 비교 파일 등록 → 본문 읽기·규칙 요약 → 엑셀 받기까지 확인한다.
CHROMIUM_PATH 가 없으면 건너뛴다.
"""
import json
import os
import re
import shutil
import ssl
import subprocess
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import pytest
from openpyxl import Workbook, load_workbook
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parent.parent
EXT_DIR = ROOT / "seat-extension"
CHROMIUM = os.environ.get("CHROMIUM_PATH")
needs_chromium = pytest.mark.skipif(not CHROMIUM, reason="CHROMIUM_PATH 미설정")


@pytest.mark.parametrize("name", ["cafe-common.js", "template.js", "bg-render.js"])
def test_shared_files_match(name):
    """댓글 수집기(chrome-extension)와 함께 쓰는 파일은 두 확장에서 같아야 한다 (한쪽만 고치지 않도록)."""
    assert (EXT_DIR / name).read_bytes() == (ROOT / "chrome-extension" / name).read_bytes(), name


def test_manifest():
    m = json.loads((EXT_DIR / "manifest.json").read_text())
    assert m["name"] == "LX3 / ME1 / MX5 전동 시트 작동 위험 상황 관련 게시물 수집"
    other = json.loads((ROOT / "chrome-extension" / "manifest.json").read_text())
    assert m["key"] != other["key"]  # 댓글 수집기와 다른 확장 (따로 설치·따로 저장)
    for f in [*m["icons"].values(), m["background"]["service_worker"], *m["content_scripts"][0]["js"]]:
        assert (EXT_DIR / f).exists(), f

CAFES = {"palicafe": "50001", "evcafe": "50002", "hardcafe": "50003"}  # 영문 주소 → 카페 번호
# 첫 화면에 카페 번호 단서가 없는 카페 (hardcafe 는 글 화면 주소로 알아내고, nocafe 는 끝내 못 알아냄)
NO_HINT_HOME = {"hardcafe", "nocafe"}
# (카페 번호, 글번호) → (제목, 목록 날짜, 본문)
POSTS = {
    ("50001", "1001"): ("3열 탈 때 2열 시트에 아이 발 끼임", "2025.03.10.",
                        "어제 3열에 아이가 타다가 2열 전동 시트가 복귀하면서 발이 끼였어요.\n센서가 있는데도 멈추지 않아 깜짝 놀랐습니다."),
    ("50001", "1002"): ("3열 시트 가죽 질감 좋네요", "2025.04.01.", "3열 시트 가죽이 생각보다 좋습니다."),
    ("50001", "1003"): ("2열 시트 폴딩 조심하세요", "2024.12.01.", "양산 전 글"),
    ("50002", "2001"): ("아이오닉9 3열 위험하네요", "2025.02.02.",
                        "아이오닉9 3열에서 내릴 때 2열 시트 원터치 폴딩하면 손이 낄 뻔했습니다. 다들 조심하세요."),
    ("50002", "2002"): ("아이오닉5 2열 끼임", "2025.02.03.", "아이오닉5 2열 시트 이야기"),
}
# naver.me 짧은 주소 → 넘어가는 주소
SHORTS = {"abc123": "https://cafe.naver.com/evcafe/2001?art=xyz", "notcafe": "https://blog.naver.com/someone/1"}
# 검색어 → 결과에 나오는 글번호 (카페마다)
SEARCH = {
    ("50001", "3열 시트"): ["1002", "1001", "1003"],
    ("50001", "시트 끼임"): ["1001"],
    ("50002", "아이오닉9 3열"): ["2002", "2001"],
}


def search_html(cafe, q, page):
    ids = SEARCH.get((cafe, q), []) if page == 1 else []
    rows = "".join(
        f'<li class="article-item"><a class="article_title" href="/f-e/cafes/{cafe}/articles/{aid}"><strong>{POSTS[(cafe, aid)][0]}</strong></a>'
        f'<span class="date">{POSTS[(cafe, aid)][1]}</span></li>'
        for aid in ids
    ) or '<p class="empty">검색 결과가 없습니다.</p>'
    return f"<html><head><title>검색 : 네이버 카페</title></head><body><ul>{rows}</ul></body></html>"


def start_server(tmp_path: Path):
    cert, key = tmp_path / "cert.pem", tmp_path / "key.pem"
    subprocess.run(["openssl", "req", "-x509", "-newkey", "rsa:2048", "-nodes", "-days", "1", "-subj", "/CN=cafe.naver.com",
                    "-keyout", str(key), "-out", str(cert)], check=True, capture_output=True)
    requests = []
    short_hits = []

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            requests.append(self.path)
            u = urlparse(self.path)
            if self.headers.get("Host", "").startswith("naver.me"):  # 짧은 주소 → 카페 글 주소로 넘김
                short_hits.append((self.path, self.headers.get("Cookie")))
                target = SHORTS.get(u.path.strip("/"))
                self.send_response(302 if target else 404)
                if target:
                    self.send_header("Location", target)
                self.send_header("Content-Length", "0")
                self.end_headers()
                return
            body = "<html><body>없는 화면</body></html>"
            m = re.match(r"/f-e/cafes/(\d+)/menus/0$", u.path)
            if m:
                qs = parse_qs(u.query)
                body = search_html(m.group(1), qs["q"][0], int(qs.get("page", ["1"])[0]))
            elif re.match(r"^/([a-z]+)/?$", u.path):  # 카페 첫 화면: 화면 속 링크에 카페 번호
                slug = u.path.strip("/")
                body = ("<html><body>카페 소개</body></html>" if slug in NO_HINT_HOME
                        else f'<html><body><a href="/f-e/cafes/{CAFES[slug]}/menus/0">전체글보기</a></body></html>')
            elif (m := re.match(r"^/([a-z]+)/(\d+)$", u.path)) and m.group(1) == "hardcafe":  # 짧은 글 주소 → 새 카페 글 주소로
                self.send_response(302)
                self.send_header("Location", f"/f-e/cafes/50003/articles/{m.group(2)}")
                self.end_headers()
                return
            elif re.match(r"^/f-e/cafes/50003/articles/\d+$", u.path):
                body = '<html><body><h3 class="title_text">하드 카페 글</h3><div class="se-main-container"><p>본문</p></div></body></html>'
            elif m := re.match(r"^/([a-z]+)/(\d+)$", u.path):
                cafe = CAFES[m.group(1)]
                title, date, text = POSTS[(cafe, m.group(2))]
                body = (f'<html><body><div class="ArticleTitle"><h3 class="title_text">{title}</h3></div>'
                        f'<div class="article_info"><span class="date">{date} 14:05</span></div>'
                        f'<div class="se-main-container">{"".join(f"<p>{line}</p>" for line in text.split(chr(10)))}</div></body></html>')
            data = body.encode()
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    tls = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
    tls.load_cert_chain(cert, key)
    server.socket = tls.wrap_socket(server.socket, server_side=True)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    server.requests = requests
    server.short_hits = short_hits
    return server


def make_compare_files(tmp_path: Path):
    """사용자 일일 온라인 동향 리스트 · 로데이터와 같은 머리글의 가짜 파일."""
    trend = Workbook()
    ws = trend.active
    ws.title = "전체"
    ws.append(["차종", "재원", "불만구분", "매체", "동호회명", "게시일", "게시글 위치", "내용", "중복", "작성자"])
    ws.append(["LX3", "-", "품질(주의/경고/심각)", "동호회", "팰리 카페", "2025-03-10", "https://cafe.naver.com/palicafe/1001", "끼임", "신규", "-"])
    ws.append(["MX5", "-", "품질", "동호회", "다른 카페", "2025-06-30", "https://cafe.naver.com/other/1", "", "신규", "-"])
    ws.append(["ME1", "-", "품질", "동호회", "전기차", "2025-01-02", "https://cafe.naver.com/evcafe/1999", "", "신규", "-"])
    trend.save(tmp_path / "251231_일일동향.xlsx")
    raw = Workbook()
    ws = raw.active
    ws.title = "현대자동차"
    ws.append([None, "커뮤니티명", "차종", "게시일", "구분", "게시글 위치", "주요 내용"])
    ws.append([1, "팰리 카페", "LX3", "2025-03-10", "품질", "https://cafe.naver.com/palicafe/1001?art=abc", "끼임"])
    ws.append([2, "전기차", "ME1", "2025-02-02", "품질", "https://naver.me/abc123", "3열"])  # 짧은 주소로만 있는 글
    ws.append([3, "전기차", "ME1", "2025-02-03", "품질", "https://naver.me/notcafe", "블로그"])
    ws.append([5, "전기차", "ME1", "2025-02-04", "품질", "https://naver.me/expired", "만료된 짧은 주소"])  # 넘어가지 않음(404)
    ws.append([4, "다른 동호회", "NE", "2025-02-03", "품질", "https://naver.me/othercar", "다른 차종"])  # 우리 동호회·차종 아님 → 풀지 않음
    raw.save(tmp_path / "251231_raw_data.xlsx")
    return [tmp_path / "251231_일일동향.xlsx", tmp_path / "251231_raw_data.xlsx"]


def open_panels(page):
    """화면 준비가 끝난 뒤 접혀 있는 2·3번 칸(비교 파일 · Claude AI 설정)을 펼친다."""
    page.locator("#login", has_text="네이버").wait_for()
    page.evaluate("document.querySelectorAll('details.panel').forEach(d => d.open = true)")


@needs_chromium
def test_seat(tmp_path: Path):
    # 실제 전달용 비교 자료(compare-data.json, 저장소에 없음) 대신 작은 가짜 자료를 넣은 복사본으로 시험
    ext_dir = tmp_path / "ext"
    shutil.copytree(EXT_DIR, ext_dir, ignore=shutil.ignore_patterns("compare-data.json"))
    (ext_dir / "compare-data.json").write_text(json.dumps({"version": "t1", "files": [
        {"id": "b1", "name": "미리 받은 동향.xlsx", "kind": "trend", "keys": ["evcafe/1999"], "slugs": {"evcafe": 1},
         "minDate": None, "maxDate": None, "rows": 1},
        {"id": "b2", "name": "미리 받은 로데이터.xlsx", "kind": "raw", "keys": ["hardcafe/3001"], "slugs": {"hardcafe": 1},
         "minDate": None, "maxDate": None, "rows": 1}]}, ensure_ascii=False))
    server = start_server(tmp_path)
    port = server.server_address[1]
    with sync_playwright() as p:
        ctx = p.chromium.launch_persistent_context(
            str(tmp_path / "profile"), executable_path=CHROMIUM, headless=True, timezone_id="Asia/Seoul",
            args=[f"--disable-extensions-except={ext_dir}", f"--load-extension={ext_dir}",
                  f"--host-resolver-rules=MAP cafe.naver.com 127.0.0.1:{port}, MAP naver.me 127.0.0.1:{port}", "--no-proxy-server", "--ignore-certificate-errors"],
            accept_downloads=True,
        )
        naver_api = []
        ctx.route("https://apis.naver.com/**", lambda route: (naver_api.append(route.request.url), route.fulfill(status=500, body="")))
        sw = ctx.service_workers[0] if ctx.service_workers else ctx.wait_for_event("serviceworker")
        ext = f"chrome-extension://{sw.url.split('/')[2]}"
        page = ctx.new_page()
        errors = []
        page.on("pageerror", lambda e: errors.append(str(e)))
        ctx.add_init_script("window.__seatAutoConfirm = true;")  # 화면 안 확인 상자는 자동으로 '확인'
        page.goto(f"{ext}/seat.html")
        page.evaluate("CAFE_ARTICLE_DELAY[0] = CAFE_ARTICLE_DELAY[1] = CAFE_PAGE_DELAY[0] = CAFE_PAGE_DELAY[1] = 10")
        # 2·3번 칸: 준비된 칸(비교 파일 등록됨)은 접고, 준비가 덜 된 칸(API 키 없음)은 펼친 채 · 접혀도 제목 옆에 상태
        page.locator("#login", has_text="네이버").wait_for()
        assert not page.evaluate("document.getElementById('panelCompare').open") and page.evaluate("document.getElementById('panelAi').open")
        assert "파일 2개 (동향 1 · 로데이터 1) · 짧은 주소 모두 확인" in page.inner_text("#compareState")
        assert "API 키 없음" in page.inner_text("#aiState")
        open_panels(page)

        assert page.inner_text("h1") == "LX3 / ME1 / MX5 전동 시트 작동 위험 상황 관련 게시물 수집"
        assert page.evaluate("getComputedStyle(document.querySelector('h1')).color") == "rgb(0, 0, 0)"  # 제목 글자 검정
        assert page.locator("nav.menu").count() == 0  # 별도 프로그램: 다른 화면 메뉴 없음
        # 함께 받은 비교 자료는 처음 열 때 자동 등록
        assert page.locator("#compareTable tbody tr").count() == 2 and "미리 등록" in page.inner_text("#compareTable")
        assert page.get_attribute(".brand .logo", "src") == "icons/logo-seat.png"
        assert page.evaluate("document.querySelector('.brand .logo').naturalWidth") > 0
        # 위쪽 띠: 에르메스 오렌지
        assert page.evaluate("getComputedStyle(document.querySelector('.topbar')).backgroundColor") == "rgb(243, 112, 33)"
        # AI 키: 결과가 없어도 바로 보이고, 저장하면 AI 가 켜지고, 다시 열어도 남아 있다
        ctx.route("https://api.anthropic.com/v1/models/**", lambda route: route.fulfill(
            status=200 if route.request.headers.get("x-api-key") == "sk-ant-good" else 401,
            headers={"content-type": "application/json", "access-control-allow-origin": "*"},
            body=json.dumps({"id": "claude-opus-5-5", "type": "model", "display_name": "Opus", "created_at": "2026-01-01T00:00:00Z"})
            if route.request.headers.get("x-api-key") == "sk-ant-good" else json.dumps({"type": "error", "error": {"type": "authentication_error", "message": "invalid x-api-key"}})))
        assert page.is_visible("#aiKey") and not page.is_checked("#aiOn")
        page.fill("#aiKey", "sk-ant-bad")
        page.click("#aiTest")
        page.locator("#aiStatus", has_text="API 키가 올바르지 않습니다").wait_for()
        page.fill("#aiKey", "sk-ant-good")
        page.click("#aiTest")
        page.locator("#aiStatus", has_text="키가 맞습니다").wait_for()
        assert page.is_checked("#aiOn")
        page.reload()
        page.evaluate("CAFE_ARTICLE_DELAY[0] = CAFE_ARTICLE_DELAY[1] = CAFE_PAGE_DELAY[0] = CAFE_PAGE_DELAY[1] = 10")
        open_panels(page)
        assert "✔ 키 저장됨 · Claude Opus 5.5 사용" in page.inner_text("#aiState")
        page.locator("#aiStatus", has_text="저장된 키로").wait_for()
        assert page.input_value("#aiKey") == "sk-ant-good" and page.is_checked("#aiOn") and "저장된 키로" in page.inner_text("#aiStatus")
        page.uncheck("#aiOn")  # 아래는 규칙 요약으로 시험
        page.locator("#aiOn").dispatch_event("change")

        # 기본 동호회 12곳, 차종 3개(시작일: 양산 시점)
        assert len(page.input_value("#cafeList").split("\n")) == 12
        assert page.evaluate("CARS.map(c => carSettings()[c.code].start)") == ["2024-12-20", "2024-12-16", "2023-08-15"]

        # 테스트용 동호회 두 곳 · 검색어 2개 · MX5 는 찾지 않음
        page.locator("#settings summary").click()
        page.fill("#cafeList", "LX3 | 팰리 카페 | palicafe\nME1 | 전기차 | evcafe | 혼합\nLX3 | 하드 카페 | hardcafe\nLX3 | 없는 카페 | nocafe")
        page.locator("#cafeList").dispatch_event("change")
        page.fill("#queries", "3열 시트, 시트 끼임")
        page.locator("#queries").dispatch_event("change")
        page.fill("#mixedQueries", "{차종} 3열")
        page.locator("#mixedQueries").dispatch_event("change")
        page.locator("#carRows tr[data-car=MX5] .carOn").uncheck()
        # 꼭 넣을 글: 기본은 요청서 예시 2개 · 시험에서는 단어 판단으로 빠지는 글(1002)과 양산 시점 이전 글(1003)
        assert page.input_value("#mustList") == "https://cafe.naver.com/gamewinds/4299497\nhttps://cafe.naver.com/tobloomer/133964"
        page.fill("#mustList", "https://cafe.naver.com/palicafe/1002\nhttps://cafe.naver.com/palicafe/1003")
        page.locator("#mustList").dispatch_event("change")
        assert "동호회·검색어 7개" in page.inner_text("#jobInfo")
        # 카페 번호 칸: 카페 주소를 붙여 넣으면 번호만 남는다
        slug_inputs = page.locator("#cafeInfo tbody input")
        slug_inputs.nth(0).fill("https://cafe.naver.com/f-e/cafes/50001/menus/0?viewType=L")
        slug_inputs.nth(0).dispatch_event("change")
        assert slug_inputs.nth(0).input_value() == "50001"

        page.click("#search")
        try:
            page.locator("#summary", has_text="새로 찾은 글").wait_for(timeout=180000)
        except Exception:
            raise AssertionError(f"{page.inner_text('#summary')}\n{page.inner_text('#log')}\n{errors}\n{server.requests}")
        assert not errors, errors
        rows = page.locator("#results tbody tr")
        got = {rows.nth(i).locator("a").inner_text(): rows.nth(i).locator("input").is_checked() for i in range(rows.count())}
        assert got == {
            "3열 탈 때 2열 시트에 아이 발 끼임": True,
            "3열 시트 가죽 질감 좋네요": True,  # 위험 단어가 없어 자동 체크는 안 되지만 꼭 넣을 글
            "아이오닉9 3열 위험하네요": True,
            "아이오닉5 2열 끼임": False,  # 혼합 동호회인데 차종 이름 없음
            "2열 시트 폴딩 조심하세요": True,  # 양산 시점 이전이라 검색에선 빠지지만 꼭 넣을 글 → 글을 열어 직접 넣음
        }, (got, page.inner_text("#log"))
        assert "새로 찾은 글 4개" in page.inner_text("#summary")
        assert "꼭 넣을 글 2개 중 검색으로 찾음 1개" in page.inner_text("#summary")
        log = page.inner_text("#log")
        assert "palicafe/1002: 검색엔 나왔지만 단어 판단으로 체크 안 됨" in log and "palicafe/1003: 검색으로 못 찾아 직접 넣음 (양산 시점 이전 글)" in log, log
        must = page.evaluate("items.filter(x => x.must).map(x => [x.key, x.car, x.bodyRead, x.body])")
        assert sorted(must) == [["palicafe/1002", "LX3", False, ""], ["palicafe/1003", "LX3", True, "양산 전 글"]], must
        # 아래 엑셀 확인을 단순하게: 꼭 넣을 글 두 개는 지운다
        for t in ["3열 시트 가죽 질감 좋네요", "2열 시트 폴딩 조심하세요"]:
            rows.filter(has_text=t).locator(".del").click()
        rows = page.locator("#results tbody tr")
        assert rows.count() == 3  # 검색으로 찾은 4개 + 직접 넣은 1개 - 지운 2개
        # 카페 번호: 첫 화면에서, 첫 화면에 없으면 비교 파일에 있는 그 카페 글 화면에서 알아냄
        assert page.evaluate("[clubIds.palicafe, clubIds.evcafe, clubIds.hardcafe]") == ["50001", "50002", "50003"]
        # 끝내 못 알아낸 카페는 한 번만 알리고 나머지 검색어는 건너뜀 · 진단 파일 받기 가능
        log = page.inner_text("#log")
        assert log.count("✖ 없는 카페(nocafe)") == 1 and "'카페 번호' 칸에 붙여 넣은" in log, log
        assert page.is_visible("#diag")
        assert not naver_api

        # 비교 파일: 머리글로 종류 자동 판단, f-e 주소·?art= 붙은 주소도 같은 글로
        page.set_input_files("#compareFiles", [str(f) for f in make_compare_files(tmp_path)])
        page.locator("#compareInfo", has_text="2개 파일을 등록").wait_for()
        kinds = page.locator("#compareTable tbody select")
        assert sorted(kinds.nth(i).input_value() for i in range(4)) == ["raw", "raw", "trend", "trend"]
        # naver.me 짧은 주소: 풀기 전에는 N, 풀면 같은 글로 잡힘 (우리 동호회·차종 행만 · 로그인 정보 없이 · 따라가지 않음)
        page.locator("#shortText", has_text="3개를 아직 풀지 않았습니다").wait_for()
        assert rows.filter(has_text="아이오닉9 3열 위험하네요").locator("td.yn").inner_text() == "N / N"
        page.click("#resolveShorts")
        try:
            page.locator("#shortText", has_text="모두 확인했습니다").wait_for(timeout=30000)
        except Exception:
            raise AssertionError(f"{page.inner_text('#shortText')} / {page.inner_text('#summary')} / {server.short_hits} / {page.evaluate('shortMap')}")
        assert "열리지 않는 주소 2개" in page.inner_text("#shortText")
        assert "포함 여부(Y/N)가 바뀐 글 1개 (그중 체크한 글 1개)" in page.inner_text("#summary"), page.inner_text("#summary")  # 블로그로 넘어감 1 + 만료 1 → 다시 묻지 않음
        assert page.is_hidden("#resolveShorts")
        assert sorted(p for p, _ in server.short_hits) == ["/abc123", "/expired", "/notcafe"], server.short_hits
        assert all(c is None for _, c in server.short_hits)
        assert not any("/evcafe/2001" in r for r in server.requests)  # 짧은 주소가 넘어가는 카페 글은 열지 않음
        yn = {rows.nth(i).locator("a").inner_text(): rows.nth(i).locator("td.yn").inner_text() for i in range(rows.count())}
        assert yn["3열 탈 때 2열 시트에 아이 발 끼임"] == "Y / Y" and yn["아이오닉9 3열 위험하네요"] == "N / Y", yn

        # 본문 읽기 + 규칙 요약 (AI 끔)
        page.click("#readBodies")
        page.locator("#summary", has_text="규칙으로 요약했습니다").wait_for(timeout=120000)
        assert not errors, errors
        texts = page.evaluate("items.filter(x => x.checked).map(x => [x.title, x.bodyRead, x.summary, x.reason])")
        by = {t[0]: t for t in texts}
        assert by["3열 탈 때 2열 시트에 아이 발 끼임"][1] and "끼였어요" in by["3열 탈 때 2열 시트에 아이 발 끼임"][2]
        assert by["3열 탈 때 2열 시트에 아이 발 끼임"][3] == ""  # 온라인 동향에 있으면 누락 사유 없음
        assert by["아이오닉9 3열 위험하네요"][3].startswith("로데이터에는 수집되었으나") and by["아이오닉9 3열 위험하네요"][3].endswith("추정")

        # 다시 열어도 결과·비교 파일 유지
        page.reload()
        page.evaluate("CAFE_ARTICLE_DELAY[0] = CAFE_ARTICLE_DELAY[1] = 10")
        open_panels(page)
        page.locator("#results tbody tr").nth(2).wait_for()
        page.locator("#compareTable tbody tr").nth(3).wait_for()
        assert page.locator("#results tbody tr").count() == 3 and page.locator("#compareTable tbody tr").count() == 4

        # 작업 중에는 뒤로가기·백스페이스 막기
        page.evaluate("showBusy('테스트', 0, 1)")
        assert page.evaluate("navGuardOn") is True
        page.keyboard.press("Backspace")
        page.go_back()
        assert page.url.endswith("seat.html") and page.locator("#guardNotice").count() == 1
        page.evaluate("hideBusy()")
        assert page.evaluate("navGuardOn") is False

        with page.expect_download() as dl:
            page.click("#saveList")
        path = tmp_path / dl.value.suggested_filename
        dl.value.save_as(path)
        assert re.match(r"^\d{8}_LX3·ME1·MX5 전동 시트 작동 위험 상황 관련 게시물_1차 검수\.xlsx$", path.name), path.name
        page.locator("#saveInfo", has_text="데이터 DB(1차 검수)에 오늘 파일로 저장").wait_for()
        assert "엑셀에 넣은 글 2건" in page.inner_text("#diffs")

        # 전체 raw 데이터: 찾은 글 전체(체크 안 된 글 포함) · 판단 근거·본문 · DB 에 따로 보관
        with page.expect_download() as dl:
            page.click("#saveRaw")
        raw_path = tmp_path / dl.value.suggested_filename
        dl.value.save_as(raw_path)
        assert raw_path.name.endswith("_전체 raw data.xlsx"), raw_path.name
        page.locator("#saveInfo", has_text="데이터 DB(로우 엑셀)에 오늘 파일로 저장").wait_for()
        rws = load_workbook(raw_path).active
        assert rws["A1"].font.name == "현대산스 Text" and rws["E2"].font.name == "현대산스 Text"
        head = [c.value for c in rws[1]]
        assert head[:8] == ["순번", "차종", "동호회명", "게시일", "제목", "URL", "최종 포함(체크)", "온라인동향/로데이터 포함 여부"] and head[-1] == "본문"
        body_rows = [[c.value for c in r] for r in rws.iter_rows(min_row=2)]
        assert len(body_rows) == 3 and sum(1 for r in body_rows if r[6] == "O") == 2  # 찾은 글 3개(체크 안 된 글 포함) 중 체크 2개
        by_title = {r[4]: r for r in body_rows}
        assert by_title["아이오닉5 2열 끼임"][6] is None and by_title["아이오닉9 3열 위험하네요"][7] == "N / Y"
        assert "끼였어요" in by_title["3열 탈 때 2열 시트에 아이 발 끼임"][-1]

        # 데이터 DB: 받은 엑셀이 오늘 날짜로 보관됨 · 받기 · 수정본 올리기 · 작업 데이터 백업/불러오기
        page.click("a.tool[href='seat-db.html']")
        page.wait_for_url("**/seat-db.html")
        page.locator("section[data-kind=seatRaw] tbody tr .chip").first.wait_for()
        raw_chips = page.locator("section[data-kind=seatRaw] tbody tr").first.locator(".chip").all_inner_texts()
        assert raw_chips == ["찾은 글 3", "체크 2"], raw_chips
        db_rows = page.locator("section[data-kind=seat] tbody tr")
        db_rows.first.locator(".counts").wait_for()
        assert db_rows.count() == 1 and "최종" in db_rows.first.inner_text()
        assert db_rows.first.locator(".chip").all_inner_texts() == ["글 2", "LX3 1", "ME1 1", "MX5 0", "N/N 0", "N/Y 1", "Y/Y 1"]
        assert "온라인 동향 N" in db_rows.first.locator(".counts").inner_text()
        assert db_rows.first.locator(".pill").inner_text() == "최종" and "첫 기록" in db_rows.first.locator(".counts").inner_text()
        assert "찾은 글" in page.inner_text("#work") and "3개" in page.inner_text("#work")
        assert "아직 기록 없음" in page.inner_text("#work")  # 규칙 요약만 써서 AI 비용 기록 없음
        with page.expect_file_chooser() as fc:
            db_rows.first.locator(".up").click()
        fc.value.set_files(str(path))
        page.locator("#dbInfo", has_text="바꿨습니다").wait_for()
        assert "직접 올림" in db_rows.first.inner_text()
        with page.expect_download() as dl:
            page.click("#backup")
        backup = tmp_path / dl.value.suggested_filename
        dl.value.save_as(backup)
        data = json.loads(backup.read_text())
        assert data["app"] == "seat-risk" and len(data["data"]["seatItems"]) == 3
        assert "aiKey" not in json.dumps(data) and "sk-ant-good" not in backup.read_text()  # API 키는 백업에 넣지 않음
        assert all(not f.get("builtin") for f in data["data"]["seatCompare"]) and len(data["data"]["seatCompare"]) == 2
        page.evaluate("chrome.storage.local.set({ seatItems: [] })")
        page.set_input_files("#restoreFile", str(backup))
        page.locator("#dbInfo", has_text="백업을 불러왔습니다").wait_for()
        assert page.evaluate("chrome.storage.local.get('seatItems').then(s => s.seatItems.length)") == 3
        assert page.evaluate("chrome.storage.local.get('seatCompare').then(s => s.seatCompare.length)") == 4
        assert not errors, errors
        ctx.close()
    server.shutdown()

    wb = load_workbook(path, rich_text=True)
    ws = wb.active
    assert [c.value for c in ws[3]][1:] == ["순번", "차종", "동호회명", "게시일", "제목", "URL", "주요내용 요약", "온라인동향/로데이터 포함 여부", "미포함시 누락 사유", "주제 관련도"]
    r4 = [c.value for c in ws[4]]
    assert r4[1:4] == [1, "LX3", "팰리 카페"] and r4[5] == "3열 탈 때 2열 시트에 아이 발 끼임"
    assert ws["E4"].value.strftime("%Y-%m-%d") == "2025-03-10" and ws["E4"].number_format == "yyyy-mm-dd"
    assert r4[6] == "https://cafe.naver.com/palicafe/1001" and ws["G4"].hyperlink.target == r4[6]
    assert r4[9] is None
    r5 = [c.value for c in ws[5]]
    assert r5[1:4] == [2, "ME1", "전기차"] and r5[6] == "https://cafe.naver.com/evcafe/2001"
    assert r5[9].startswith("로데이터에는 수집되었으나")
    # 포함 여부: 한 칸 안에서 Y 파랑 · N 빨강
    assert ws["F4"].font.name == "현대산스 Text" and ws["B3"].font.name == "현대산스 Text"  # 엑셀 글꼴
    cell = ws["I5"].value
    assert str(cell) == "N / Y"
    colors = [(b.text, b.font.color.rgb if b.font and b.font.color else None) for b in cell if hasattr(b, "font")]
    assert colors == [("N", "FFE00000"), (" / ", None), ("Y", "FF0070C0")], colors
    assert all(b.font.rFont == "현대산스 Text" for b in cell if hasattr(b, "font"))
    assert str(ws["I4"].value) == "Y / Y"
    assert ws.max_row == 5


systems = []  # 최종 판단 요청의 지시문 (담당자 추가 규칙이 붙는지 확인)
reviews = []  # 최종 검수 요청의 지시문


def fake_claude(route, calls):
    """Claude API 흉내: 1차 검토(제목·미리보기)와 최종 판단(본문)을 구분해 정해진 규칙대로 답한다 (스트리밍 형식)."""
    body = json.loads(route.request.post_data)
    items = json.loads(body["messages"][0]["content"].split("\n\n", 1)[1])
    screen = "1차 검토자" in body["system"]
    calls.append("screen" if screen else "final")
    review = "최종 검수자" in body["system"]
    if review:
        calls[-1] = "review"
    elif not screen:
        systems.append(body["system"])
    if review:  # 최종 검수: 본문에 '끼' 가 있으면 관련도 95(포함), 없으면 40(제외)
        out = {"items": [{"id": x["id"], "seatIssue": "트렁크" not in x["본문"], "accident": "끼" in x["본문"], "concern": False, "partialMention": False,
                          "relevance": 95 if "끼" in x["본문"] else 40, "why": "판독 내용 기준"} for x in items]}
        reviews.append(body["system"])
    elif screen:
        out = {"items": [{"id": x["id"], "maybe": "아이" in x["제목"] + x["미리보기"], "why": "아이 위험 언급"} for x in items]}
    else:
        out = {"items": [{
            "id": x["id"], "relevant": "끼" in x["본문"], "situation": "아차 사고" if "끼" in x["본문"] else "관련 없음",
            # 여러 글을 한 번에 보낼 때 '시트커버' 글만 형식이 틀린 답(빈 요약)을 준다 → 하나씩 다시 보내면 제대로
            "summary": "" if x["제목"] == "시트커버 팝니다" and len(items) > 1 else f"{x['제목']} 관련 내용", "reason": "" if x["온라인동향포함"] else "아차 상황으로 보고 기준에 미달한 것으로 추정",
        } for x in items]}
    ev = lambda t, d: f"event: {t}\ndata: {json.dumps(d, ensure_ascii=False)}\n\n"
    sse = (ev("message_start", {"type": "message_start", "message": {"id": "m", "type": "message", "role": "assistant", "model": body["model"],
                                                                     "content": [], "stop_reason": None, "usage": {"input_tokens": 1000, "output_tokens": 0}}})
           + ev("content_block_start", {"type": "content_block_start", "index": 0, "content_block": {"type": "text", "text": ""}})
           + ev("content_block_delta", {"type": "content_block_delta", "index": 0, "delta": {"type": "text_delta", "text": json.dumps(out, ensure_ascii=False)}})
           + ev("content_block_stop", {"type": "content_block_stop", "index": 0})
           + ev("message_delta", {"type": "message_delta", "delta": {"stop_reason": "end_turn"}, "usage": {"output_tokens": 200}})
           + ev("message_stop", {"type": "message_stop"}))
    route.fulfill(status=200, headers={"content-type": "text/event-stream", "access-control-allow-origin": "*"}, body=sse)


@needs_chromium
def test_ai_review_unchecked(tmp_path: Path):
    """체크 안 된 글도 AI 가 본다: ① 제목·미리보기 1차 검토로 관련 가능한 글을 체크 ② 본문을 읽은 뒤 AI 가 체크를 정함."""
    ext_dir = tmp_path / "ext"
    shutil.copytree(EXT_DIR, ext_dir, ignore=shutil.ignore_patterns("compare-data.json"))
    server = start_server(tmp_path)
    with sync_playwright() as p:
        ctx = p.chromium.launch_persistent_context(
            str(tmp_path / "profile"), executable_path=CHROMIUM, headless=True,
            args=[f"--disable-extensions-except={ext_dir}", f"--load-extension={ext_dir}", "--no-proxy-server", "--ignore-certificate-errors",
                  f"--host-resolver-rules=MAP cafe.naver.com 127.0.0.1:{server.server_address[1]}"])
        calls = []
        ctx.route("https://api.anthropic.com/v1/messages**", lambda route: fake_claude(route, calls))
        sw = ctx.service_workers[0] if ctx.service_workers else ctx.wait_for_event("serviceworker")
        page = ctx.new_page()
        errors = []
        page.on("pageerror", lambda e: errors.append(str(e)))
        ctx.add_init_script("window.__seatAutoConfirm = true;")
        asked = lambda: page.evaluate("globalThis.__seatAsked || []")
        page.goto(f"chrome-extension://{sw.url.split('/')[2]}/seat.html")
        open_panels(page)  # 화면 준비(저장된 결과 불러오기)가 끝난 뒤
        base = "car:'LX3',cafeIndex:0,cafeName:'팰리 카페',slug:'palicafe',date:Date.UTC(2025,2,1),queries:['q'],touched:false,summary:'',reason:'',by:'',bodyNote:''"
        page.evaluate(f"""items = [
          {{key:'palicafe/1',articleId:'1',title:'큰일날 뻔했어요',text:'3열 내리다 아이가',checked:false,bodyRead:true,body:'3열에서 내리다 아이 손이 2열에 끼일 뻔',{base}}},
          {{key:'palicafe/2',articleId:'2',title:'시트커버 팝니다',text:'새 제품',checked:false,bodyRead:true,body:'시트커버 판매',{base}}},
          {{key:'palicafe/3',articleId:'3',title:'2열 끼임 사고',text:'2열 끼임',checked:true,bodyRead:true,body:'2열 폴딩 중 끼임',{base}}},
          {{key:'palicafe/4',articleId:'4',title:'아이랑 3열 사진',text:'아이 사진',checked:false,bodyRead:true,body:'3열에 앉은 아이 사진입니다',{base}}},
        ]; renderResults();""")
        page.fill("#aiKey", "sk-ant-test")
        page.click("#aiSave")
        assert page.is_checked("#aiOn") and page.is_checked("#screenOn")
        # 💰 = AI 비용이 드는 버튼 (AI 켜짐: 본문 읽기·다시 쓰기에도 붙음)
        assert page.inner_text("#readBodies").startswith("💰") and page.inner_text("#rewrite").startswith("💰")
        assert page.inner_text("#saveFinal").startswith("💰") and page.inner_text("#redoReview").startswith("💰")
        page.uncheck("#aiOn")
        page.locator("#aiOn").dispatch_event("change")
        assert not page.inner_text("#readBodies").startswith("💰") and not page.inner_text("#rewrite").startswith("💰")
        assert not page.inner_text("#saveFinal").startswith("💰")  # AI 를 끄면 검수를 하지 않으므로 비용 없음
        assert page.inner_text("#redoReview").startswith("💰")  # 검수 다시 하기는 늘 AI
        page.check("#aiOn")
        page.locator("#aiOn").dispatch_event("change")

        # ① 1차 검토 → ② (본문은 이미 읽음) 최종 판단
        page.click("#readBodies")
        page.locator("#summary", has_text="Claude AI 로 요약했습니다").wait_for(timeout=60000)
        assert any("체크 안 된 글 3개를 AI 가 제목·미리보기로" in m for m in asked()), asked()
        assert calls[0] == "screen" and "final" in calls
        got = page.evaluate("Object.fromEntries(items.map(x => [x.key, [x.checked, x.aiScreen ? x.aiScreen.maybe : null, x.by]]))")
        assert got == {
            "palicafe/1": [True, True, "ai"],  # 제목엔 단어가 없지만 1차 검토로 체크 → 본문 보고 AI 가 관련 있다고 판단
            "palicafe/2": [False, False, ""],  # 1차 검토에서 무관
            "palicafe/3": [True, None, "ai"],  # 원래 체크
            "palicafe/4": [False, True, "ai"],  # 1차 검토로 체크됐지만 본문을 보고 AI 가 무관하다고 판단 → 해제
        }, got
        assert "AI 1차(제목·미리보기): 관련 가능" in page.locator("#results tbody tr", has_text="큰일날 뻔했어요").inner_text()

        # 전체 정독: 체크 안 된 글도 본문을 읽은 것은 AI 가 판단 (1차 검토 없이)
        page.evaluate("""items.push({key:'palicafe/5',articleId:'5',title:'강아지',text:'강아지',checked:false,bodyRead:true,
          body:'2열 폴딩에 강아지 발이 끼임',car:'LX3',cafeIndex:0,cafeName:'팰리 카페',slug:'palicafe',date:Date.UTC(2025,2,2),
          queries:['q'],touched:false,summary:'',reason:'',by:'',bodyNote:''}); renderResults();""")
        page.check("#readAll")
        n_screen = calls.count("screen")
        page.click("#readBodies")
        try:
            page.wait_for_function("() => items.find(x => x.key === 'palicafe/5').by === 'ai'", timeout=20000)
        except Exception:
            raise AssertionError(f"{page.inner_text('#summary')} | {asked()[-2:]} | {calls} | {errors} | {page.evaluate('items.map(x => [x.key, x.checked, x.by, x.touched])')}")
        assert calls.count("screen") == n_screen  # 전체 정독이면 1차 검토는 건너뜀
        assert page.evaluate("items.find(x => x.key === 'palicafe/2').by") == "ai"
        # 담당자 추가 규칙(메모)이 지시문에 붙고, 모니터링 시작 전 글(팰리세이드 오너스 클럽 '26.3.11 이전)은 정해진 누락 사유
        assert "[담당자 추가 규칙 - 위 규칙보다 우선]" in systems[-1] and "tobloomer" in systems[-1]
        page.evaluate("""items.push(
          {key:'tobloomer/1',articleId:'1',car:'LX3',cafeIndex:3,cafeName:'팰리세이드 오너스 클럽',slug:'tobloomer',date:new Date(2026,1,20).getTime(),
           title:'2열 끼임',text:'2열 끼임',checked:true,bodyRead:true,body:'2열 폴딩에 끼임',queries:['q'],touched:false,summary:'',reason:'',by:'',bodyNote:''},
          {key:'tobloomer/2',articleId:'2',car:'LX3',cafeIndex:3,cafeName:'팰리세이드 오너스 클럽',slug:'tobloomer',date:new Date(2026,2,11).getTime(),
           title:'3열 끼임',text:'3열 끼임',checked:true,bodyRead:true,body:'3열 복귀에 끼임',queries:['q'],touched:false,summary:'',reason:'',by:'',bodyNote:''}); renderResults();""")
        page.click("#rewrite")
        page.wait_for_function("() => items.find(x => x.key === 'tobloomer/2').by === 'ai'", timeout=60000)
        reasons = page.evaluate("[items.find(x => x.key === 'tobloomer/1').reason, items.find(x => x.key === 'tobloomer/2').reason]")
        assert reasons[0] == "해당 동호회(팰리세이드 오너스 클럽)는 당시 모니터링 대상 동호회가 아니었음 ('26년 3월 12일부터 모니터링 대상 동호회로 지정)", reasons
        assert reasons[1] == "아차 상황으로 보고 기준에 미달한 것으로 추정", reasons  # '26.3.11 게시글부터는 모니터링 대상
        assert page.evaluate("ruleReason(items.find(x => x.key === 'tobloomer/1'), inclusion(items.find(x => x.key === 'tobloomer/1')))") == reasons[0]

        # 전차종 로데이터 작성('24.4.15) 이전 글: 누락 사유 앞에 로데이터가 없던 시기라는 내용 + 글 내용에 따른 사유
        page.evaluate("""items.push({key:'palicafe/240201',articleId:'240201',car:'MX5',cafeIndex:0,cafeName:'팰리 카페',slug:'palicafe',
          date:new Date(2024,1,1).getTime(),title:'2열 끼임 주의',text:'2열 끼임',checked:true,bodyRead:true,body:'2열 폴딩에 손이 끼일 뻔',
          queries:['q'],touched:false,summary:'',reason:'',by:'',bodyNote:''}); renderResults();""")
        page.click("#rewrite")
        page.wait_for_function("() => items.find(x => x.key === 'palicafe/240201').by === 'ai'", timeout=60000)
        r = page.evaluate("finalReason(items.find(x => x.key === 'palicafe/240201'))")
        assert r == "전차종 로데이터는 '24년 4월 15일부터 작성되어 이 글의 게시 시점에는 로데이터 파일이 없었음 · 아차 상황으로 보고 기준에 미달한 것으로 추정", r
        assert page.locator("#results tbody tr", has_text="2열 끼임 주의").locator("td").nth(7).inner_text() == r
        assert page.evaluate("finalReason(items.find(x => x.key === 'palicafe/1'))").startswith("아차 상황")  # '25년 글엔 붙지 않음
        # 로데이터에 있는 글(Y)에는 붙이지 않음 · 첫 로데이터에 담긴 4.12(금)~ 글에도 붙이지 않음
        assert page.evaluate("rawStartNote(items.find(x => x.key === 'palicafe/240201'), {trend:false, raw:true, trendCafe:true, trendRange:true})") == ""
        assert page.evaluate("rawStartNote({date:new Date(2024,3,12).getTime()}, {trend:false, raw:false})") == ""
        assert page.evaluate("rawStartNote({date:new Date(2024,3,11).getTime()}, {trend:false, raw:false})").startswith("전차종 로데이터는 '24년 4월 15일부터")

        # 실제 쓴 AI 비용: 요청마다 사용량으로 계산해 글마다 나눠 붙이고 전체 누적도 남김
        costs = page.evaluate("items.map(x => x.aiCost || 0)")
        assert all(c > 0 for c in costs[:2]), costs
        total = page.evaluate("aiSpend.won")
        assert abs(sum(costs) - total) < 0.01 and total > 0, (costs, total)
        assert page.evaluate("fmtWon(1400)") == "약 1,400원 ($1.00)"
        assert re.search(r"사용한 비용 약 [\d,]+원 \(\$\d+\.\d{2}\)", page.inner_text("#summary")), page.inner_text("#summary")
        # 비용 기록 전에 AI 로 요약한 글(기록 없음)은 추정치로 따로
        est = page.evaluate("estimateUntrackedWon([{by:'ai', body:'가'.repeat(1000), title:'t'}, {by:'ai', aiCost: 5}], 'claude-opus-5-5')")
        assert est["posts"] == 1 and est["won"] > 0, est
        assert page.evaluate("items.find(x => x.key === 'palicafe/5').checked") is True

        # 아직 본문을 안 읽은 체크한 글: 읽는 대로 AI 가 판단·요약 (다 읽을 때까지 기다리지 않음)
        page.uncheck("#readAll")
        page.evaluate("CAFE_ARTICLE_DELAY[0] = CAFE_ARTICLE_DELAY[1] = 10")
        page.evaluate(f"""items.push({{key:'palicafe/1001',articleId:'1001',title:'목록 제목',text:'3열',checked:true,bodyRead:false,body:'',{base}}}); renderResults();""")
        page.evaluate("globalThis.__seatAsked = []")
        page.click("#readBodies")
        page.wait_for_function("() => items.find(x => x.key === 'palicafe/1001').by === 'ai'", timeout=60000)
        assert any("읽는 대로 Claude AI 가 판단·요약합니다" in m for m in asked()), asked()
        assert page.evaluate("items.find(x => x.key === 'palicafe/1001').summary") == "3열 탈 때 2열 시트에 아이 발 끼임 관련 내용"
        page.locator("#summary", has_text="AI 로 판단·요약했습니다").wait_for()

        # 포함 여부로 골라 보기·정렬 (N / N · N / Y · Y / N · Y / Y)
        page.evaluate("compare = [{id:'t', name:'동향', kind:'trend', keys:['palicafe/3'], keySet:new Set(['palicafe/3']), slugs:{palicafe:1}, minDate:null, maxDate:null},"
                      "{id:'r', name:'로데이터', kind:'raw', keys:['palicafe/3','palicafe/1'], keySet:new Set(['palicafe/3','palicafe/1']), slugs:{palicafe:2}, minDate:null, maxDate:null}]; renderResults();")
        assert page.locator("#viewFilter option[value='yn:YN']").count() == 0  # Y / N 은 나올 수 없는 조합이라 없음
        page.select_option("#viewFilter", "yn:YY")
        assert [r for r in page.locator("#results tbody tr td.yn").all_inner_texts()] == ["Y / Y"]
        page.select_option("#viewFilter", "yn:NY")
        assert page.locator("#results tbody tr td.yn").all_inner_texts() == ["N / Y"]
        page.select_option("#viewFilter", "")
        page.select_option("#sortBy", "yn")
        yns = page.locator("#results tbody tr td.yn").all_inner_texts()
        assert yns == sorted(yns, key=["N / N", "N / Y", "Y / N", "Y / Y"].index) and yns[0] == "N / N" and yns[-1] == "Y / Y", yns
        assert "Y / Y 1" in page.inner_text("#ynCounts") and "N / Y 1" in page.inner_text("#ynCounts")
        # 표 머리글 '동향 / 로데이터' 는 한 줄
        lines = page.locator("#results th.nowrap").evaluate(
            "e => { const r = document.createRange(); r.selectNodeContents(e); return new Set([...r.getClientRects()].map(c => Math.round(c.top))).size; }")
        assert lines == 1, lines

        # 예전에 저장해 둔 AI 추가 규칙 메모에도 새 기본 규칙(로데이터 작성 시작일)이 한 번 덧붙는다
        page.evaluate("""chrome.storage.local.remove('seatRuleBlocks').then(() => chrome.storage.local.get('seatSettings')).then(({ seatSettings }) =>
          chrome.storage.local.set({ seatSettings: { ...seatSettings, aiRules: '## 담당자 추가 규칙\\n- 내가 적은 규칙' } }))""")
        page.evaluate("chrome.storage.local.set({ saveAsk: true })")
        page.reload()
        open_panels(page)
        rules = page.input_value("#aiRules")
        # 규칙을 덧붙여 저장할 때 API 키·저장 위치·차종 설정이 지워지지 않음 (다시 열어도 그대로)
        page.reload()
        open_panels(page)
        assert page.input_value("#aiKey") == "sk-ant-test" and page.is_checked("#saveAsk")
        page.evaluate("chrome.storage.local.set({ saveAsk: false })")
        assert page.evaluate("chrome.storage.local.get('seatCars').then(s => Object.keys(s.seatCars || {}).length)") == 3
        assert "내가 적은 규칙" in rules and "전차종 로데이터 엑셀은 '24년 4월 15일부터" in rules and "tobloomer" not in rules, rules
        page.reload()
        open_panels(page)
        assert page.input_value("#aiRules").count("전차종 로데이터 엑셀은") == 1  # 두 번 덧붙지 않음

        # AI 최종 검수 → ② 1차 검수(체크한 글 전체 + 관련도 %) · ③ 최종 파일(최종 검수 통과한 글만)
        page.evaluate("""items.push({key:'palicafe/9',articleId:'9',car:'LX3',cafeIndex:0,cafeName:'팰리 카페',slug:'palicafe',
          date:new Date(2025,5,1).getTime(),title:'3열 시트 가죽 이야기',text:'3열',checked:true,touched:true,bodyRead:true,body:'3열 시트 가죽 질감 이야기',
          queries:['q'],summary:'3열 시트 가죽 질감을 소개하는 내용',reason:'',by:'ai',bodyNote:''},
          {key:'palicafe/10',articleId:'10',car:'LX3',cafeIndex:0,cafeName:'팰리 카페',slug:'palicafe',
          date:new Date(2025,5,2).getTime(),title:'전동 트렁크 손 끼임',text:'트렁크',checked:true,touched:true,bodyRead:true,body:'전동 트렁크 닫히다 손 끼임',
          queries:['q'],summary:'전동 트렁크에 손이 끼였다는 내용',reason:'',by:'ai',bodyNote:''},
          {key:'palicafe/11',articleId:'11',car:'LX3',cafeIndex:0,cafeName:'팰리 카페',slug:'palicafe',
          date:new Date(2025,5,3).getTime(),title:'예전 기준 검수 글',text:'2열',checked:true,touched:true,bodyRead:true,body:'2열 복귀에 발 끼임',
          queries:['q'],summary:'2열 복귀 끼임',reason:'',by:'ai',bodyNote:'',review:{relevance:50,keep:false,why:'예전'}}); renderResults();""")
        n_checked = page.evaluate("items.filter(x => x.checked).length")
        # ② 1차 검수를 먼저 눌러도 검수(관련도 + 4항목)를 그 자리에서 하고 받음 · 검수할 글이 남아 있을 때만 💰
        assert page.inner_text("#saveList").startswith("💰") and page.inner_text("#saveFinal").startswith("💰")
        with page.expect_download() as dl:
            page.click("#saveList")
        first_path = tmp_path / dl.value.suggested_filename
        dl.value.save_as(first_path)
        assert first_path.name.endswith("_1차 검수.xlsx") and calls.count("review") >= 1, (first_path.name, calls)
        assert any("검수합니다" in m for m in asked())
        assert not page.inner_text("#saveList").startswith("💰") and not page.inner_text("#saveFinal").startswith("💰")
        rv = page.evaluate("Object.fromEntries(items.filter(x => x.checked).map(x => [x.key, [x.review.relevance, x.review.keep]]))")
        assert rv["palicafe/9"] == [40, False] and rv["palicafe/1"] == [95, True], rv
        assert rv["palicafe/10"][1] is False  # 전동 트렁크 이슈는 끼임이어도 배제 (①)
        assert rv["palicafe/11"] == [95, True]  # 이전 기준으로 검수한 글은 새 기준으로 다시 검수
        assert "4가지를 하나씩 다시 검수" in reviews[-1] and "[담당자 추가 규칙" in reviews[-1] and "전동 트렁크(테일게이트)" in reviews[-1]
        sws = load_workbook(first_path, rich_text=True).active
        rows1 = {sws.cell(r, 6).value: sws.cell(r, 11).value for r in range(4, sws.max_row + 1)}
        assert len(rows1) == n_checked and rows1["3열 시트 가죽 이야기"] == 0.4  # 숫자(백분율 서식) - '텍스트로 저장된 숫자' 경고 없음
        assert "최종 검수: 관련도 40% · 제외 · ①시트작동O ②사고X ③우려X ④일부언급X" in page.locator("#results tbody tr", has_text="3열 시트 가죽 이야기").inner_text()

        # ③ 최종 파일: ② 에서 한 검수를 그대로 씀 (AI 를 다시 부르지 않음)
        n_review = calls.count("review")
        with page.expect_download() as dl:
            page.click("#saveFinal")
        final_path = tmp_path / dl.value.suggested_filename
        dl.value.save_as(final_path)
        assert final_path.name.endswith("_최종.xlsx") and calls.count("review") == n_review
        fws = load_workbook(final_path, rich_text=True).active
        final_titles = [fws.cell(r, 6).value for r in range(4, fws.max_row + 1)]
        assert "3열 시트 가죽 이야기" not in final_titles and "전동 트렁크 손 끼임" not in final_titles and len(final_titles) == n_checked - 2, final_titles
        assert fws["K4"].value == 0.95 and fws["K4"].number_format == "0%"
        # 최종 파일의 '판단 기준' 시트: 관련도 구간별 기준 · 최종 검수 4항목 · 수집 기간·채널 · 포함 여부·누락 사유을 그때 설정값으로 채움 · 보고용이라 'AI' 표현 없음
        fwb = load_workbook(final_path)
        assert fwb.sheetnames == ["시트 위험상황 게시물", "판단 기준"] and fwb.active.title == "시트 위험상황 게시물"
        cws = fwb["판단 기준"]
        texts = [str(c.value) for row in cws.iter_rows() for c in row if c.value is not None]
        joined = "\n".join(texts)
        assert not any("{{" in t for t in texts), [t for t in texts if "{{" in t]
        assert "AI" not in joined, [t for t in texts if "AI" in t]
        assert not any(w in joined for w in ["단계별 건수", "작성 정보", "검색어", "비교 파일", "정렬", "꼭 넣을"])
        assert any(c.value == 1 and c.number_format == "0%" for row in cws.iter_rows() for c in row), "100% 는 숫자 셀"
        for must in ["85~95%", "70~84%", "40~69%", "10~39%", "0~9%", "① 전동 시트 작동 이슈", "팰리세이드 피플(cafe.naver.com/allumpc)", "4. 포함 여부(온라인동향 / 로데이터) · 누락 사유 기준", "모니터링 대상 동호회로 지정", "전차종 로데이터는 2024-04-15 부터 작성"]:
            assert must in joined, must
        # 색은 조건부 서식: 70% 이상 빨강 · 미만 파랑
        cf = {str(r.sqref): [(x.operator, x.formula, x.dxf.font.color.rgb) for x in r.rules] for r in fws.conditional_formatting}
        assert cf == {"K4:K5000": [("greaterThanOrEqual", ["0.7"], "FFE00000"), ("between", ["0", "0.6999"], "FF0070C0")]}, cf
        assert page.inner_text("#saveInfo").startswith(f"✔ 1차 검수 {n_checked}개 중 {n_checked - 2}개 최종 확정 (제외 2개)"), page.inner_text("#saveInfo")
        assert page.inner_text("#summary") == f"✔ 1차 검수 {n_checked}개 중 {n_checked - 2}개 최종 확정 (제외 2개)"
        page.select_option("#viewFilter", "dropped")
        assert page.locator("#results tbody tr").count() == 2
        page.select_option("#viewFilter", "")

        # 엑셀 저장 위치: '폴더 선택'으로 고른 폴더에 바로 저장 (시험에서는 폴더 창 대신 브라우저 안 임시 폴더)
        page.evaluate("document.getElementById('settings').open = true")
        assert page.inner_text("#folderName") == "크롬 다운로드 폴더" and page.is_hidden("#clearFolder")
        page.evaluate("window.showDirectoryPicker = () => navigator.storage.getDirectory()")
        page.click("#pickFolder")
        page.locator("#folderName.on").wait_for()
        assert page.is_visible("#clearFolder")
        page.click("#saveList")
        page.locator("#saveInfo", has_text="폴더에").wait_for(timeout=30000)
        saved = page.evaluate("""async () => { const d = await navigator.storage.getDirectory(); const out = [];
          for await (const [name, h] of d.entries()) if (h.kind === 'file') out.push([name, (await h.getFile()).size]); return out; }""")
        assert any(n.endswith("_LX3·ME1·MX5 전동 시트 작동 위험 상황 관련 게시물_1차 검수.xlsx") and size > 1000 for n, size in saved), saved
        page.reload()
        open_panels(page)
        page.evaluate("document.getElementById('settings').open = true")
        page.locator("#folderName.on").wait_for()  # 고른 폴더는 다시 열어도 기억
        page.click("#clearFolder")
        page.locator("#folderName", has_text="크롬 다운로드 폴더").wait_for()

        # 확인 상자는 화면 안에 떠서, 다른 탭을 보고 있어도 막히지 않고 답할 때까지 기다린다 (크롬 confirm() 은 뒤쪽 탭에서 막힘)
        other = ctx.new_page()
        other.goto("about:blank")
        other.bring_to_front()
        page.evaluate("() => { globalThis.__seatAutoConfirm = false; globalThis.__answer = null; askConfirm('계속할까요?').then(v => (globalThis.__answer = v)); }")
        assert page.is_visible(".ask-back") and page.title().startswith("❓ 확인 필요")
        assert page.evaluate("globalThis.__answer") is None  # 답하기 전에는 기다림
        page.bring_to_front()
        page.click(".ask-ok")
        assert page.evaluate("globalThis.__answer") is True and not page.title().startswith("❓")
        assert not errors, errors
        ctx.close()
    server.shutdown()
