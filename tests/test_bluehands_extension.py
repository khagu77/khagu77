"""블루핸즈/하이테크센터 정비 예약 확장프로그램 통합 테스트.

cafe.naver.com · search.naver.com 을 로컬 HTTPS 서버로 돌려
기사목록 파일 → 작성자 확인(직접 / 검색 경유 / 패스) → 지난 보고서에 이어 쓰기 → DB 목록까지 확인한다.
확장 로딩에는 일반 Chromium 이 필요하므로 CHROMIUM_PATH 가 없으면 건너뛴다.
"""
import filecmp
import os
import shutil
import ssl
import subprocess
import threading
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlparse
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
EXT_DIR = ROOT / "bluehands-extension"
CHROMIUM = os.environ.get("CHROMIUM_PATH")


def test_shared_files_match_comment_collector():
    """두 확장이 같이 쓰는 파일은 내용이 같아야 한다 (한쪽만 고치지 않도록)."""
    for name in ("template.js", "cafe-common.js"):
        assert filecmp.cmp(EXT_DIR / name, ROOT / "chrome-extension" / name, shallow=False), name


def article_html(nick, body="본문 내용입니다. 하이테크센터 예약했어요."):
    return (
        "<html><head><title>테스트 카페 : 네이버 카페</title></head><body><div class='ArticleContentBox'>"
        "<div class='article_header'><h3 class='title_text'>글 제목</h3>"
        f"<div class='WriterInfo'><div class='nick_box'><button class='nickname'>{nick}</button></div></div></div>"
        f"<div class='article_container'><div class='se-main-container'><p>{body}</p></div></div>"
        "</div></body></html>"
    )


# 지난주 iroid/5520064 처럼 글 주소로 들어가면 알림창으로 삭제를 알리고 카페 첫 화면으로 보내는 경우
DELETED_ALERT = "<html><body><script>alert('삭제되었거나 존재하지 않는 게시글입니다.'); location.replace('/aacafe');</script></body></html>"
DELETED = "<html><body><div class='error_content'><p>삭제되었거나 존재하지 않는 게시글입니다.</p></div></body></html>"
MEMBERS_ONLY = "<html><body><div class='error_content'><p>카페 멤버만 볼 수 있습니다.</p><a>카페 가입하기</a></div></body></html>"

# 제목 → 검색 결과에 나올 링크
SEARCH_RESULTS = {
    "블루핸즈 예약 문의": ["https://cafe.naver.com/aacafe/9999?art=WRONG", "https://cafe.naver.com/aacafe/102?art=TOKEN102"],
    "넥쏘 하이테크센터 후기": ["https://cafe.naver.com/nexocafe/201?art=TOKEN201"],
}


def start_fake_naver(tmp_path: Path):
    cert, key = tmp_path / "cert.pem", tmp_path / "key.pem"
    subprocess.run(
        ["openssl", "req", "-x509", "-newkey", "rsa:2048", "-nodes", "-days", "1", "-subj", "/CN=cafe.naver.com",
         "-keyout", str(key), "-out", str(cert)],
        check=True, capture_output=True,
    )
    hits = []

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            host = self.headers.get("Host", "").split(":")[0]
            u = urlparse(self.path)
            q = parse_qs(u.query)
            hits.append((host, self.path, self.headers.get("Referer", "")))
            if host == "search.naver.com":
                query = unquote(q.get("query", [""])[0]).strip('"')
                links = SEARCH_RESULTS.get(query, [])
                body = "<html><body><div id='main_pack'>" + (
                    "".join(f"<div class='total_wrap'><a class='title_link' href='{h}'>{query}</a></div>" for h in links)
                    if links else "<p>'검색결과가 없습니다'</p>"
                ) + "</div></body></html>"
            elif u.path == "/aacafe/101":
                body = article_html("울산ll테스터")
            elif u.path == "/aacafe/102":
                body = article_html("검색ll작성자") if q.get("art") == ["TOKEN102"] else MEMBERS_ONLY
            elif u.path == "/nexocafe/201":
                body = article_html("넥쏘ll회원") if q.get("art") == ["TOKEN201"] else MEMBERS_ONLY
            elif u.path == "/aacafe/104":  # 삭제된 게시물 (알림창)
                body = DELETED_ALERT
            elif u.path == "/nexocafe/202":  # 삭제된 게시물 (화면 안내)
                body = DELETED
            else:  # aacafe/103 등: 언제나 막힘
                body = MEMBERS_ONLY
            data = body.encode()
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            # 네이버처럼 '틀 안에서 열지 마세요' 머리글 - 확장이 자기 화면 안의 틀에 한해 빼야 읽힌다
            self.send_header("X-Frame-Options", "SAMEORIGIN")
            self.send_header("Content-Security-Policy", "frame-ancestors 'self'")
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    tls = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
    tls.load_cert_chain(cert, key)
    server.socket = tls.wrap_socket(server.socket, server_side=True)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    server.hits = hits
    return server


def make_list_xls(path: Path):
    """모니터링 사이트의 기사목록 .xls (실제로는 HTML 표) 흉내 - 최신 글이 위."""
    rows = [
        ("디올뉴 넥쏘카페 / 수소차 동호회", "넥쏘 하이테크센터 후기", "2026-10-11 09:00:00", "https://cafe.naver.com/nexocafe/201"),
        ("디올뉴 넥쏘카페 / 수소차 동호회", "지워진 넥쏘 글", "2026-10-10 20:00:00", "https://cafe.naver.com/nexocafe/202"),
        ("AA 동호회", "못 보는 글", "2026-10-10 10:00:00", "https://cafe.naver.com/aacafe/103"),
        ("AA 동호회", "지워진 글", "2026-10-09 18:00:00", "https://cafe.naver.com/aacafe/104"),
        ("AA 동호회", "블루핸즈 예약 문의", "2026-10-09 15:30:00", "https://cafe.naver.com/aacafe/102"),
        ("AA 동호회", "하이테크센터 정비 예약했어요", "2026-10-08 11:25:00", "https://cafe.naver.com/aacafe/101"),
        ("AA 동호회", "이미 보고한 글", "2026-10-06 08:00:00", "https://cafe.naver.com/aacafe/100"),
        ("뉴스", "블로그 글", "2026-10-06 07:00:00", "https://blog.naver.com/someone/1"),
    ]
    trs = "".join(
        f"<tr height=24><td align=center>{i + 1}</td><td align=center>CAFE</td><td align=center>{cafe}</td>"
        f"<td align=center>네이버</td><td>{title}</td><td align=center>-</td><td align=center></td><td>{when}</td>"
        f"<td>{when}</td><td></td><td>{url}</td></tr>"
        for i, (cafe, title, when, url) in enumerate(rows)
    )
    path.write_text(
        '<html xmlns:o="urn:schemas-microsoft-com:office:office"><head><meta http-equiv="Content-Type" content="text/html; charset=utf-8">'
        "</head><body><table border=1><tr height=28><th>No</th><th>채널</th><th>출처</th><th>포털</th><th>제목</th><th>분류</th>"
        f"<th>성향</th><th>작성시간</th><th>수집시간</th><th>댓글수</th><th>Link</th></tr>{trs}</table></body></html>",
        encoding="utf-8",
    )


def make_report(path: Path):
    """지난 보고서: 양식에 2행(마지막 연번 5)."""
    shutil.copy(EXT_DIR / "report-template.xlsx", path)
    wb = load_workbook(path)
    ws = wb.active
    styles = [c._style for c in ws[2]]
    ws.delete_rows(2)
    for r, row in enumerate([
        [4, "동호회", "AA 동호회", datetime(2026, 10, 5, 9, 0), "지난주 글", "서울ll누구", "https://cafe.naver.com/aacafe/99"],
        [5, "동호회", "AA 동호회", datetime(2026, 10, 6, 8, 0), "이미 보고한 글", "부산ll누구", "https://cafe.naver.com/aacafe/100"],
    ], start=2):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(r, c, v)
            cell._style = styles[c - 1]
        ws.row_dimensions[r].height = 24
    wb.save(path)


@pytest.mark.skipif(not CHROMIUM, reason="CHROMIUM_PATH 미설정")
def test_bluehands(tmp_path: Path):
    from playwright.sync_api import sync_playwright

    server = start_fake_naver(tmp_path)
    port = server.server_address[1]
    list_xls = tmp_path / "기사목록20261012.xls"
    make_list_xls(list_xls)
    report = tmp_path / "20261005_정비 예약 관련 테스트 보고서.xlsx"
    make_report(report)

    with sync_playwright() as p:
        ctx = p.chromium.launch_persistent_context(
            str(tmp_path / "profile"),
            executable_path=CHROMIUM,
            headless=True,
            timezone_id="Asia/Seoul",
            args=[
                f"--disable-extensions-except={EXT_DIR}",
                f"--load-extension={EXT_DIR}",
                f"--host-resolver-rules=MAP cafe.naver.com 127.0.0.1:{port}, MAP search.naver.com 127.0.0.1:{port}",
                "--no-proxy-server",
                "--ignore-certificate-errors",
            ],
            accept_downloads=True,
        )
        api_calls = []
        ctx.route("https://apis.naver.com/**", lambda route: (api_calls.append(route.request.url), route.fulfill(status=500, body="")))
        sw = ctx.service_workers[0] if ctx.service_workers else ctx.wait_for_event("serviceworker")
        ext = f"chrome-extension://{sw.url.split('/')[2]}"
        dialogs, leave_prompts = [], []

        def on_dialog(d):
            if d.type == "beforeunload":  # 작업 중 새로고침·닫기 확인 → 테스트에서는 '나가기'
                leave_prompts.append(1)
                d.accept()
            else:
                dialogs.append(d.message)
                d.dismiss()

        ctx.on("page", lambda pg: pg.on("dialog", on_dialog))
        page = ctx.new_page()
        page.goto(f"{ext}/bluehands.html")
        assert page.inner_text("h1") == "블루핸즈/하이테크센터 정비 예약 관련 수집 시스템"

        # 보고 일정: 월요일, 월요일이 쉬는 날이면 다음 근무일
        sched = page.evaluate("""() => [
            ['2026-10-12', '2026-10-05', '2026-10-07', '2027-02-10', '2026-10-13'].map(s => {
                const [y, m, d] = s.split('-').map(Number);
                return ymd(reportDayOfWeek(new Date(y, m - 1, d)).day);
            }),
            ymd(reportDayOfWeek(new Date(2026, 9, 19), parseExtraHolidays('2026-10-19 회사 휴무')).day),
            reportSchedule(new Date(2026, 9, 7)).next.day.getDate(),
        ]""")
        assert sched == [["2026-10-12", "2026-10-06", "2026-10-06", "2027-02-10", "2026-10-12"], "2026-10-20", 12]

        # 테스트는 빨리: 글 사이 기다림을 줄인다
        page.evaluate("CAFE_ARTICLE_DELAY.splice(0, 2, 10, 20); CAFE_PAGE_DELAY.splice(0, 2, 10, 20)")

        page.set_input_files("#listFile", str(list_xls))
        page.locator("#listInfo", has_text="✔").wait_for()
        assert "글 8개 (2026-10-06 07:00 ~ 2026-10-11 09:00)" in page.inner_text("#listInfo")
        page.set_input_files("#reportFile", str(report))
        page.locator("#reportInfo", has_text="✔").wait_for()
        assert "2행(마지막 연번 5)" in page.inner_text("#reportInfo")
        rows = page.locator("#items tbody tr")
        assert rows.count() == 8
        assert "이미 보고서에 있음" in rows.nth(1).inner_text()  # 날짜순: 블로그(07:00) → 이미 보고한 글(08:00) → …
        assert "네이버 카페 글 주소가 아님" in rows.nth(0).inner_text()

        pages_before = len(ctx.pages)
        page.click("#start")
        # 작업 중 표시: 맨 위 로딩 바 · 탭 제목
        page.locator("#busy").wait_for(state="visible")
        assert "작성자 확인 중" in page.inner_text("#busy")
        rows.nth(2).locator("td", has_text="✔ 확인").wait_for(timeout=120000)
        assert page.title().startswith("⏳ ")
        # 뒤로 가기(마우스 제스처 등)는 막힌다 - 화면·작업 그대로
        page.go_back()  # 브라우저의 진짜 뒤로 가기 (마우스 제스처·Alt+←와 같음)
        page.locator("#toast", has_text="뒤로 가기가 막혀").wait_for()
        assert page.url.endswith("/bluehands.html") and page.locator("#busy").is_visible()
        # 실수로 새로고침(창 닫기와 같음) → 확인 창이 뜨고, 나간 뒤 다시 열면 '이어서 진행'
        page.reload()
        assert leave_prompts, "작업 중 나가기 확인 창이 떠야 함"
        page.evaluate("CAFE_ARTICLE_DELAY.splice(0, 2, 10, 20); CAFE_PAGE_DELAY.splice(0, 2, 10, 20)")
        page.locator("#resumeBox", has_text="중간에 끊겼습니다").wait_for()
        assert "지난번에 고른 파일" in page.inner_text("#reportInfo")  # 보고서 파일도 기억
        assert rows.nth(2).locator("td").nth(6).inner_text() == "✔ 확인"
        page.click("#resume")
        page.locator("#summary", has_text="작성자 확인을 마쳤습니다").wait_for(timeout=300000)
        assert page.locator("#busy").is_hidden() and not page.title().startswith("⏳")
        # 이미 끝난 글은 다시 열지 않는다
        assert len([h for h in server.hits if h[1].startswith("/aacafe/101")]) == 1
        # 탭·창을 하나도 띄우지 않는다
        assert len(ctx.pages) == pages_before, [pg.url for pg in ctx.pages]
        assert "최소화된 창" not in page.inner_text("#log")  # 보이지 않는 틀에서 다 읽음
        # 삭제된 게시물 2개(알림창 1 · 화면 안내 1, 넥쏘카페 글은 검색에서 못 찾아 글 주소로 확인)는
        # 목록에 빨간 취소선으로 남는다
        assert rows.count() == 8
        deleted = page.locator("#items tbody tr.deleted")
        assert deleted.count() == 2
        assert [deleted.nth(i).locator("td").nth(3).inner_text() for i in range(2)] == ["지워진 글", "지워진 넥쏘 글"]
        assert "삭제된 게시물" in deleted.nth(0).locator("td").nth(6).inner_text()
        assert deleted.nth(0).locator("td").nth(3).evaluate("e => getComputedStyle(e).textDecorationLine") == "line-through"
        assert "삭제된 게시물 2개(엑셀에서 뺌)" in page.inner_text("#itemCount")
        # 이미 보고서에 있는 글은 보고서의 작성자를 보여 준다
        assert rows.nth(1).locator("td").nth(5).inner_text() == "부산ll누구"
        live = page.locator("#items tbody tr:not(.deleted)")
        got = [
            (live.nth(i).locator("td").nth(0).inner_text(), live.nth(i).locator("td").nth(3).inner_text(),
             live.nth(i).locator(".nick").inner_text() if live.nth(i).locator(".nick").count()
             else live.nth(i).locator("input.author").input_value(),
             live.nth(i).locator("td").nth(6).inner_text())
            for i in range(2, 6)
        ]
        assert got == [
            ("6", "하이테크센터 정비 예약했어요", "울산ll테스터", "✔ 확인"),
            ("7", "블루핸즈 예약 문의", "검색ll작성자", "✔ 확인(검색 경유)"),
            ("8", "못 보는 글", "", "패스 - 검색에서 같은 글을 찾지 못함"),
            ("9", "넥쏘 하이테크센터 후기", "넥쏘ll회원", "✔ 확인(검색 경유)"),
        ], got
        # 알림창은 화면에 뜨지 않는다(확장이 연 탭에서만 막고 기록)
        assert dialogs == [], dialogs
        assert "이어서 진행" in page.inner_text("#log")
        assert "누적 5건 · 전주 2건 대비 ▲ 3 건" in page.inner_text("#weekCompare")
        assert page.locator("#weekCompare .up").evaluate("e => getComputedStyle(e).color") == "rgb(224, 36, 36)"
        assert "확인 2개 · 패스 1개 · 삭제된 게시물 2개(엑셀에서 뺌)" in page.inner_text("#summary")

        # 넥쏘카페는 처음부터 검색으로: 글 주소를 바로 열지 않는다
        assert not [h for h in server.hits if h[1] == "/nexocafe/201"]
        # 검색 결과 링크로 들어갈 때 검색 화면을 거쳐 온 것으로 (Referer)
        art = [h for h in server.hits if h[1].startswith("/aacafe/102?art=")]
        assert art and "search.naver.com" in art[0][2]

        try:
            with page.expect_download(timeout=15000) as dl:
                page.click("#save")
        except Exception:
            raise AssertionError(page.inner_text("#saveInfo"))
        out = tmp_path / dl.value.suggested_filename
        dl.value.save_as(out)
        today = datetime.now(ZoneInfo("Asia/Seoul")).strftime("%Y%m%d")
        assert out.name == f"{today}_정비 예약 관련 테스트 보고서.xlsx"
        assert "3개 추가 (연번 6~9), 패스 1개는 넣지 않았습니다, 삭제된 게시물 2개는 넣지 않았습니다" in page.inner_text("#saveInfo")
        assert "누적 5건 · 전주 2건 대비 ▲ 3 건" in page.inner_text("#saveInfo .compare")

        # DB 목록: 지난 보고서 2행 + 새 글 3행
        page.click("nav.menu a[data-tab=db]")
        page.locator("#dbCount", has_text="DB 전체 5행").wait_for()
        page.fill("#dbQuery", "넥쏘")
        assert page.locator("#dbTable tbody tr").count() == 1
        assert page.locator("#history tbody tr").count() == 1
        with page.expect_download() as dl:
            page.click("#dbExport")
        db_file = tmp_path / dl.value.suggested_filename
        dl.value.save_as(db_file)
        assert db_file.name == f"{today}_정비 예약 관련 DB 목록_검색결과.xlsx"
        ctx.close()
    server.shutdown()
    assert api_calls == []

    # 시트 이름은 설정의 고정 이름(원본 '기사목록20260911 (1)' 대신), 문서 속성의 시트 목록도 같이
    out_wb = load_workbook(out)
    assert out_wb.sheetnames == ["정비 예약 모니터링"]
    assert load_workbook(db_file).sheetnames == ["정비 예약 모니터링"]
    import zipfile
    app = zipfile.ZipFile(out).read("docProps/app.xml").decode()
    assert "기사목록" not in app
    ws = out_wb.active
    data = [[c.value for c in row] for row in ws.iter_rows(min_row=1)]
    assert data[0] == ["No", "채널", "채널명", "작성일", "제목", "작성자", "URL"]
    assert data[1:] == [
        [4, "동호회", "AA 동호회", datetime(2026, 10, 5, 9, 0), "지난주 글", "서울ll누구", "https://cafe.naver.com/aacafe/99"],
        [5, "동호회", "AA 동호회", datetime(2026, 10, 6, 8, 0), "이미 보고한 글", "부산ll누구", "https://cafe.naver.com/aacafe/100"],
        [6, "동호회", "AA 동호회", datetime(2026, 10, 8, 11, 25), "하이테크센터 정비 예약했어요", "울산ll테스터", "https://cafe.naver.com/aacafe/101"],
        [7, "동호회", "AA 동호회", datetime(2026, 10, 9, 15, 30), "블루핸즈 예약 문의", "검색ll작성자", "https://cafe.naver.com/aacafe/102"],
        [9, "동호회", "디올뉴 넥쏘카페 / 수소차 동호회", datetime(2026, 10, 11, 9, 0), "넥쏘 하이테크센터 후기", "넥쏘ll회원", "https://cafe.naver.com/nexocafe/201"],
    ]
    # 새 행은 바로 위(마지막) 행과 같은 서식(글꼴·테두리·표시 형식·행 높이), URL 에는 링크
    for r in (4, 5, 6):
        for c in range(1, 8):
            assert ws.cell(r, c)._style == ws.cell(3, c)._style, (r, c)
        assert ws.row_dimensions[r].height == 24
        assert ws.cell(r, 7).hyperlink.target == ws.cell(r, 7).value
    assert ws.cell(4, 6).font.name == "현대산스 Text Medium" and ws.cell(4, 4).number_format == "m/d/yy h:mm"

    exported = [[c.value for c in row] for row in load_workbook(db_file).active.iter_rows(min_row=2)]
    assert exported == [data[5]]
