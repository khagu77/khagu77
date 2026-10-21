"""'이슈 모니터링 보고서 작성' 확장 통합 테스트.

실제 네이버 뉴스·유튜브·클리앙·Claude API 대신 가짜 응답(route)을 넣고,
1 목적·채널·키워드 → 2 게시글 수집(AI 관련성) → 3 댓글 수집 → 4 성향 분석 → 5 PPT·최종 DB 저장까지 확인한다.
확장 로딩에는 headless-shell 이 아닌 일반 Chromium 이 필요하므로 CHROMIUM_PATH 가 없으면 건너뛴다.
"""
import io
import json
import os
import re
import zipfile
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import pytest
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

EXT_DIR = Path(__file__).resolve().parent.parent / "issue-report-extension"
CHROMIUM = os.environ.get("CHROMIUM_PATH")
pytestmark = pytest.mark.skipif(not CHROMIUM, reason="CHROMIUM_PATH 미설정")

YT_CHANNEL = "UCabcdefghijklmnopqrstuv"

# ---------------------------------------------------------------- 네이버 뉴스
# 실제 네이버 뉴스 검색처럼: 제목·요약문 링크가 같은 언론사 주소, 링크 안에 화면낭독용 '새 창 열림', 「네이버뉴스」 링크는 따로
NEWS_SEARCH = (
    "<html><body><div class='list'>"
    '<div class="news_item"><div class="profile"><a class="press" href="https://www.yna.co.kr/">연합뉴스<span class="blind">새 창 열림</span></a>'
    '<span class="info">2026.08.06.</span><a href="https://n.news.naver.com/mnews/article/001/0000000001">네이버뉴스</a></div>'
    '<a class="news_tit" href="https://www.yna.co.kr/view/AKR0001"><span class="headline">아반떼 가격 공개 기사</span><span class="blind">새 창 열림</span></a>'
    '<a class="news_dsc" href="https://www.yna.co.kr/view/AKR0001">현대차는 최근 신형 아반떼에서 불거진 가격 논란 이후 투싼의 가격 산정을 놓고 막판까지 고심한 것으로 알려졌다. 투싼은 가격 인상분을...<span class="blind">새 창 열림</span></a></div>'
    '<div class="news_item"><a class="news_tit" href="https://n.news.naver.com/mnews/article/002/0000000002">댓글 없는 아반떼 기사 제목</a><span class="info">2026.08.06.</span></div>'
    # 언론사 자체 사이트 기사 (「네이버뉴스」 링크 없음)
    '<div class="news_item"><div class="profile"><a class="press" href="https://www.hankyung.com/">한국경제<span class="blind">새 창 열림</span></a><span class="info">2026.08.06.</span></div>'
    '<a class="news_tit" href="https://www.hankyung.com/article/2026080612345"><span class="headline">아반떼 가격 공개 외부 언론사 기사</span><span class="blind">새 창 열림</span></a>'
    '<a class="news_dsc" href="https://www.hankyung.com/article/2026080612345">기존 4세대 투싼과 비슷한 가격대를 형성하면서 앞서 출시한 디 올 뉴 아반떼와의 가격 격차가 대폭 줄어들었다. 소비자들의 평가는...<span class="blind">새 창 열림</span></a></div>'
    "</div></body></html>"
)
NEWS_COMMENTS = {
    "news001,0000000001": ["가격이 너무 비싸요 4천은 심했다", "현기 배짱장사 답이 없네", "디자인 좋네요 기대됩니다"],
    "news002,0000000002": [],
}


def news_article(oid, aid):
    return (
        f'<html><head><meta property="og:title" content="기사 {aid}"></head><body>'
        f'<div class="media_end_head_top_logo"><img title="연합뉴스"></div>'
        f'<h2 id="title_area">현대차 디 올 뉴 아반떼 가격 공개 {aid[-1]}</h2>'
        f'<span class="media_end_head_info_datestamp_time" data-date-time="2026-08-06 09:00:00">2026.08.06</span>'
        f'<article id="dic_area">아반떼 가격이 2398만원부터 시작합니다.</article></body></html>'
    )


def cbox(url):
    q = parse_qs(urlparse(url).query)
    texts = NEWS_COMMENTS.get(q["objectId"][0], [])
    data = {
        "success": True,
        "result": {
            "count": {"comment": len(texts)},
            "pageModel": {"totalPages": 1},
            "commentList": [
                {"commentNo": i + 1, "contents": t, "regTime": "2026-08-06T09:3%d:00+0900" % i, "userName": "u", "replyCount": 0}
                for i, t in enumerate(texts)
            ],
        },
    }
    return "cb(" + json.dumps(data, ensure_ascii=False) + ");"


# ---------------------------------------------------------------- 유튜브
YT_FEED = f"""<?xml version="1.0" encoding="UTF-8"?>
<feed xmlns:yt="http://www.youtube.com/xml/schemas/2015" xmlns:media="http://search.yahoo.com/mrss/" xmlns="http://www.w3.org/2005/Atom">
 <entry><yt:videoId>abcdefghijk</yt:videoId><title>신형 아반떼 가격 리뷰</title><published>2026-08-06T00:30:00+00:00</published>
  <author><name>자동차채널</name></author><media:group><media:description>CN8 가격 분석</media:description></media:group></entry>
 <entry><yt:videoId>oldoldoldol</yt:videoId><title>지난주 영상</title><published>2026-07-20T00:00:00+00:00</published>
  <author><name>자동차채널</name></author><media:group><media:description>-</media:description></media:group></entry>
</feed>"""
YT_WATCH = (
    '<html><script>ytcfg.set({"INNERTUBE_API_KEY":"KEY","INNERTUBE_CLIENT_VERSION":"2.0"});</script>'
    '<script>var ytInitialPlayerResponse = {"videoDetails":{"title":"신형 아반떼 가격 리뷰","author":"자동차채널","shortDescription":"CN8"},'
    '"microformat":{"playerMicroformatRenderer":{"publishDate":"2026-08-06T09:30:00+09:00"}}};</script>'
    "<script>var ytInitialData = "
    + json.dumps({"contents": {"itemSectionRenderer": {"sectionIdentifier": "comment-item-section", "contents": [
        {"continuationItemRenderer": {"continuationEndpoint": {"continuationCommand": {"token": "T1"}}}}]}}})
    + ";</script></html>"
)


def yt_comment(cid, text):
    return {"payload": {"commentEntityPayload": {"properties": {"commentId": cid, "content": {"content": text}, "publishedTime": "1시간 전"}, "author": {"displayName": "a"}}}}


YT_NEXT = {"frameworkUpdates": {"entityBatchUpdate": {"mutations": [yt_comment("Y1", "노조 월급 주려고 올렸네 현기"), yt_comment("Y2", "가격 비싸서 쏘나타 삽니다")]}}}

# ---------------------------------------------------------------- 클리앙
CLIEN_LIST = (
    "<html><body>"
    '<div class="list_item"><a href="/service/board/cm_car/111">아반떼 가격 진짜 비싸네요</a><span class="timestamp">2026-08-06</span></div>'
    '<div class="list_item"><a href="/service/board/cm_car/112">무관한 아반떼 가격 광고</a><span class="timestamp">2026-08-06 07:00:00</span></div>'
    '<div class="list_item"><a href="/service/board/cm_car/113">타이어 교체 후기</a><span class="timestamp">2026-08-06 06:00:00</span></div>'
    '<div class="list_item"><a href="/service/board/cm_car/100">아반떼 가격 옛날 글</a><span class="timestamp">2026-08-01 06:00:00</span></div>'
    "</body></html>"
)
CLIEN_VIEW = (
    '<html><body><div class="post_author"><span class="view_count date">2026-08-06 08:00:00</span></div>'
    '<div class="post_article">아반떼 가격 너무 올랐습니다</div>'
    '<div class="comment_row"><div class="comment_view">그래도 차는 잘 나왔어요 좋네요</div><span class="timestamp">2026-08-06 08:10:00</span></div>'
    '<div class="comment_row"><div class="comment_view">옵션 장사 심하네요 비싸</div><span class="timestamp">2026-08-06 08:20:00</span></div>'
    "</body></html>"
)


# ---------------------------------------------------------------- Claude API (스트리밍 응답 흉내)
def sse(obj_text):
    events = [
        ("message_start", {"type": "message_start", "message": {"id": "msg_1", "type": "message", "role": "assistant", "model": "claude-sonnet-5-5",
                                                                "content": [], "stop_reason": None, "stop_sequence": None, "usage": {"input_tokens": 100, "output_tokens": 1}}}),
        ("content_block_start", {"type": "content_block_start", "index": 0, "content_block": {"type": "text", "text": ""}}),
        ("content_block_delta", {"type": "content_block_delta", "index": 0, "delta": {"type": "text_delta", "text": obj_text}}),
        ("content_block_stop", {"type": "content_block_stop", "index": 0}),
        ("message_delta", {"type": "message_delta", "delta": {"stop_reason": "end_turn", "stop_sequence": None}, "usage": {"output_tokens": 50}}),
        ("message_stop", {"type": "message_stop"}),
    ]
    return "".join(f"event: {e}\ndata: {json.dumps(d, ensure_ascii=False)}\n\n" for e, d in events)


def senti_of(text):
    if "현기" in text or "노조" in text:
        return "단순비난"
    if "비싸" in text or "심했" in text or "장사" in text:
        return "부정"
    if "좋" in text or "기대" in text:
        return "긍정"
    return "중립"


def claude(body, calls):
    system = "".join(b["text"] for b in body["system"]) if isinstance(body.get("system"), list) else body.get("system", "")
    user = body["messages"][0]["content"]
    calls.append({"system": system, "user": user, "model": body["model"]})
    if "수집 키워드 추천 규칙" in system:
        assert "CN8 가격 공개 이후 가격·옵션 여론 파악" in system
        return {"lines": ["CN8|아반떼 가격|가격표", "아반떼 옵션 유료"], "exclude": ["렌트"], "note": "차종과 가격 표현을 묶었습니다."}
    if "게시물 관련성 판독" in system:
        posts = re.findall(r'<post id="(\d+)">\n.*?\n.*?\n제목: (.*?)\n', user)
        res = [{"id": i, "relevant": "무관" not in t and "타이어" not in t, "reason": "테스트", "s": senti_of(t)} for i, t in posts]
        return {"results": res}
    if "댓글 성향 분석" in system:
        rows = re.findall(r"^(\d+)\t(?:└답글 )?(.*)$", user, re.M)
        return {"results": [{"id": i, "s": senti_of(t)} for i, t in rows]}
    if "★ 댓글 화면 판독 규칙 ★" in system:
        found = re.findall(r"^댓글내용: (.*)$", user, re.M)
        return {"comments": [{"text": t, "date": "", "reply": False} for t in found]}
    if "★ 보고서 작성 규칙 ★" in system:
        ids, rows = {}, {}
        for sec in re.split(r"\n## ", user)[1:]:
            name = sec.split(" ")[0]
            ids[name] = re.findall(r"^(c\d+)\t", sec, re.M)
            rows[name] = re.findall(r"^(c\d+)\t\((?:.*?)\) (.*)$", sec, re.M)
        short = {k: ids.get(k, [])[:2] for k in ["긍정", "중립", "부정", "단순비난"]}
        long = {k: ids.get(k, []) for k in ["긍정", "중립", "부정", "단순비난"]}
        # 실제 AI 처럼 id 대신 'c5 (언론) 글'·댓글 글 자체로 돌려주거나, 아예 비워 두는 경우
        short["부정"] = [f"{i} (언론) {t}" for i, t in rows.get("부정", [])][:2]
        long["부정"] = [t for _, t in rows.get("부정", [])]
        short["긍정"] = []
        return {
            "title": "CN8 가격 공개 관련 여론 동향",
            "summary": ["테스트 요약 첫째 문장 확인 中", "테스트 요약 둘째 문장 다수 확인"],
            "quotes_short": short,
            "quotes_long": long,
        }
    return {}


def wait_js(page, expr, timeout=30000):
    """확장 화면은 eval 이 막혀 있어 wait_for_function 대신 evaluate 를 되풀이한다."""
    import time

    end = time.time() + timeout / 1000
    while time.time() < end:
        if page.evaluate(expr):
            return
        page.wait_for_timeout(200)
    raise AssertionError(f"시간 안에 참이 되지 않음: {expr}")


def test_issue_report(tmp_path: Path):
    claude_calls = []
    downloads = {}
    with sync_playwright() as p:
        ctx = p.chromium.launch_persistent_context(
            str(tmp_path / "profile"),
            executable_path=CHROMIUM,
            headless=True,
            timezone_id="Asia/Seoul",
            args=[f"--disable-extensions-except={EXT_DIR}", f"--load-extension={EXT_DIR}", "--no-proxy-server"],
            accept_downloads=True,
        )

        def naver(route):
            url = route.request.url
            if "search.naver.com" in url:
                start = parse_qs(urlparse(url).query).get("start", ["1"])[0]
                route.fulfill(status=200, content_type="text/html; charset=utf-8", body=NEWS_SEARCH if start == "1" else "<html></html>")
            elif "apis.naver.com/commentBox" in url:
                route.fulfill(status=200, content_type="text/javascript; charset=utf-8", body=cbox(url))
            elif "n.news.naver.com" in url:
                m = re.search(r"article/(\d+)/(\d+)", url)
                route.fulfill(status=200, content_type="text/html; charset=utf-8", body=news_article(m.group(1), m.group(2)))
            else:
                route.fulfill(status=200, content_type="text/html", body="<html></html>")

        def youtube(route):
            url = route.request.url
            if "/feeds/videos.xml" in url:
                route.fulfill(status=200, content_type="application/xml", body=YT_FEED)
            elif "/youtubei/v1/next" in url:
                route.fulfill(status=200, content_type="application/json", body=json.dumps(YT_NEXT))
            elif "/watch" in url:
                route.fulfill(status=200, content_type="text/html", body=YT_WATCH)
            else:
                route.fulfill(status=200, content_type="text/html", body=f'<html><meta property="og:title" content="자동차채널"><script>{{"externalId":"{YT_CHANNEL}"}}</script></html>')

        def clien(route):
            url = route.request.url
            body = CLIEN_VIEW if re.search(r"/cm_car/\d+", url) else CLIEN_LIST
            route.fulfill(status=200, content_type="text/html; charset=utf-8", body=body)

        def anthropic(route):
            body = json.loads(route.request.post_data)
            out = json.dumps(claude(body, claude_calls), ensure_ascii=False)
            route.fulfill(status=200, headers={"content-type": "text/event-stream", "access-control-allow-origin": "*"}, body=sse(out))

        ctx.route("https://*.naver.com/**", naver)
        ctx.route("https://www.youtube.com/**", youtube)
        ctx.route("https://www.clien.net/**", clien)
        ctx.route("https://api.anthropic.com/**", anthropic)

        sw = ctx.service_workers[0] if ctx.service_workers else ctx.wait_for_event("serviceworker")
        ext_id = sw.url.split("/")[2]
        page = ctx.new_page()
        errors = []
        page.on("pageerror", lambda e: errors.append(str(e)))
        dismiss_next = []
        page.on("dialog", lambda d: d.dismiss() if dismiss_next and dismiss_next.pop() else d.accept())
        page.goto(f"chrome-extension://{ext_id}/app.html")

        def save_download(button):
            with page.expect_download(timeout=60000) as dl:
                page.click(button)
            path = tmp_path / dl.value.suggested_filename
            dl.value.save_as(path)
            return path

        # 설정: API 키
        page.click('nav .tab[data-tab="settings"]')
        page.fill("#apiKey", "sk-ant-test")
        page.click("#saveSettings")
        page.locator("#settingsMsg", has_text="저장").wait_for()

        # 1. 목적·채널·키워드
        page.click('nav .tab[data-tab="setup"]')
        page.fill("#purpose", "CN8 가격 공개 이후 가격·옵션 여론 파악")
        page.fill("#issueName", "CN8 가격 공개")
        page.fill("#rangeFrom", "2026-08-05T07:00")
        page.fill("#rangeTo", "2026-08-06T10:00")
        page.click("#suggestKeywords")
        page.locator("#suggestMsg", has_text="추천 완료").wait_for(timeout=30000)
        assert page.input_value("#keywords") == "CN8|아반떼 가격|가격표\n아반떼 옵션 유료\n-렌트"
        # 조합이 많아도 둘째 줄(계약 취소)이 검색어에서 빠지지 않는다
        page.fill("#keywords", "아반떼|CN8|아방|디올뉴아반떼 투싼|NX5|디올뉴투싼 가격|가격표|비싸|역전|차이\n아반떼|CN8|아방|디올뉴아반떼 계약 취소|취소|해지|환불|위약금")
        page.dispatch_event("#keywords", "input")
        assert '"아반떼 계약 취소"' in page.inner_text("#queryPreview") and '"아반떼 계약 위약금"' in page.inner_text("#queryPreview")
        assert '"아반떼 투싼 가격"' in page.inner_text("#queryPreview")
        page.fill("#keywords", "CN8|아반떼 가격\n-렌트")
        assert '"CN8 가격"' in page.inner_text("#queryPreview") and '"아반떼 가격"' in page.inner_text("#queryPreview")
        page.uncheck("#chCafe")
        page.fill("#newsUrls", "https://n.news.naver.com/mnews/article/001/0000000001")
        # 기본 유튜브 채널 48개가 모두 선택된 채로 나온다 → 테스트는 가짜 채널 하나만
        assert page.locator("#ytList input:checked").count() == 48
        assert page.locator("#ytList label", has_text="모터그래프").count() == 1
        page.click("#ytNone")
        page.fill("#ytAdd", YT_CHANNEL)
        page.click("#ytAddBtn")
        page.locator("#ytList label", has_text="자동차채널").wait_for()
        for cb in page.locator("#commList input").all():
            if "clien" not in cb.get_attribute("value"):
                cb.uncheck()
        page.click("#goCollect")
        page.locator("#tab-posts").wait_for(state="visible")

        # 2. 게시글 수집 + AI 관련성
        page.click("#collectPosts")
        page.locator("#postsSummary", has_text="관련 게시글").wait_for(timeout=120000)
        summary = page.inner_text("#postsSummary")
        # 언론 3(댓글 있는 기사·댓글 없는 기사·언론사 자체 사이트 기사 모두) + 유튜브 1(기간 밖 영상 제외) + 클리앙 1(무관 1·키워드 없음 1·기간 밖 1 제외)
        assert "관련 게시글 5건" in summary, (summary, page.evaluate("job.posts.map(p => [p.key, p.title, p.source, p.date])"))
        # 게시글 자체 성향: AI 판독 때 함께 정해 표에 보이고, 손으로 고칠 수 있다
        clien_row = page.locator("#postsTable tbody tr", has_text="아반떼 가격 진짜 비싸네요")
        assert "시간 확인 전" in clien_row.inner_text()
        assert clien_row.locator("select").input_value() == "부정"
        assert "게시글 자체 성향 (체크한 글 5건)" in page.inner_text("#postsSentiment")
        assert "부정 1건" in page.inner_text("#postsSentiment")
        clien_row.locator("select").select_option("단순비난")
        assert "단순비난 1건" in page.inner_text("#postsSentiment")
        # 'AI 관련성 다시 판독'에서 [취소] = 체크는 그대로 두고 성향만 새로 채움 (손으로 고친 성향도 유지)
        page.locator("#postsTable tbody tr", has_text="아반떼 가격 공개 외부 언론사 기사").locator("input[type=checkbox]").uncheck()
        dismiss_next.append(True)
        page.click("#rejudge")
        wait_js(page, "document.querySelector('#postsLog').textContent.split('AI 판독 완료').length > 2", timeout=30000)
        wait_js(page, "!document.body.classList.contains('working')", timeout=30000)
        assert not page.locator("#postsTable tbody tr", has_text="아반떼 가격 공개 외부 언론사 기사").locator("input[type=checkbox]").is_checked()
        assert clien_row.locator("select").input_value() == "단순비난"
        page.locator("#postsTable tbody tr", has_text="아반떼 가격 공개 외부 언론사 기사").locator("input[type=checkbox]").check()
        sys_rel0 = next(c["system"] for c in claude_calls if "게시물 관련성 판독" in c["system"])
        assert "게시글 자체(제목·내용)가 이번 이슈 대상에 대해 보이는 성향" in sys_rel0
        posts_xlsx = save_download("#dlPosts")
        senti_posts = save_download("#dlPostsSenti")
        assert senti_posts.name.endswith("_게시물 리스트_성향분석 포함.xlsx")
        wbs = load_workbook(senti_posts)
        assert [c.value for c in wbs["커뮤니티"][1]] == ["No", "출처", "제목", "작성시간", "Link", "성향"]
        assert wbs["커뮤니티"]["F2"].value == "단순비난" and wbs["커뮤니티"]["F1"].font.name == "현대산스 Text Medium"
        cf = [(str(r.sqref), [x.text for x in r.rules]) for r in wbs["언론"].conditional_formatting]
        assert cf and cf[0][0].startswith("F2:F") and set(cf[0][1]) == {"긍정", "중립", "부정", "단순비난"}  # 성향 칸 색
        wb = load_workbook(posts_xlsx)
        assert wb.sheetnames == ["언론", "유튜브", "커뮤니티"]
        assert [c.value for c in wb["언론"][1]] == ["No", "출처", "제목", "작성시간", "Link"]
        news_rows = {r[4]: r[1] for r in wb["언론"].iter_rows(min_row=2, values_only=True)}
        news_titles = {r[4]: r[2] for r in wb["언론"].iter_rows(min_row=2, values_only=True)}
        assert news_titles["https://www.hankyung.com/article/2026080612345"] == "아반떼 가격 공개 외부 언론사 기사"  # 요약문이 아니라 제목
        assert all("새 창" not in str(v) for r in wb["언론"].iter_rows(min_row=2, values_only=True) for v in r)
        assert all(r[3] for r in wb["언론"].iter_rows(min_row=2, values_only=True))  # 작성시간이 비지 않음
        assert news_rows["https://www.hankyung.com/article/2026080612345"] == "한국경제"
        assert "https://n.news.naver.com/mnews/article/002/0000000002" in news_rows  # 댓글 없는 기사도 목록에
        assert len(news_rows) == 3
        assert wb["언론"]["B2"].value == "연합뉴스"
        assert wb["언론"]["E2"].hyperlink.target == "https://n.news.naver.com/mnews/article/001/0000000001"
        assert wb["커뮤니티"]["C2"].value == "아반떼 가격 진짜 비싸네요"
        assert wb["언론"]["A1"].font.name == "현대산스 Text Medium"

        # 3. 댓글 수집
        page.click("#goComments")
        page.click("#collectComments")
        # 작업 중에는 로딩바가 보이고, 뒤로 가기(마우스 제스처·단축키)를 막는다
        page.locator("#workBar").wait_for(state="visible")
        page.go_back()
        page.keyboard.press("Alt+ArrowLeft")
        page.wait_for_timeout(500)
        assert page.url.endswith("/app.html")
        assert "뒤로·앞으로 가기를 막았습니다" in page.inner_text("#workBarText")
        page.locator("#commentsSummary", has_text="댓글을 모은 글 5/5").wait_for(timeout=120000)
        assert "언론사 자체 사이트 — 네이버 댓글 없음" in page.inner_text("#commentsTable")
        # 목록에 날짜만 보였던 커뮤니티 글은 글 화면의 정확한 작성 시각으로 바뀜
        assert page.evaluate("job.posts.filter(p => p.key === 'clien.net:111').map(p => [p.exact, new Date(p.date).getHours()])[0]") == [True, 8]
        assert "댓글 없음 — 건너뜀" in page.inner_text("#commentsTable")
        page.locator("#workBar").wait_for(state="hidden")
        assert "댓글 7개" in page.inner_text("#commentsSummary")
        # 표에서 고른 글만 댓글 다시 수집
        page.locator("#commentsTable tbody tr", has_text="아반떼 가격 진짜 비싸네요").locator("input[type=checkbox]").check()
        assert page.inner_text("#recollectSel") == "선택한 글 댓글 다시 수집 (1)"
        page.click("#recollectSel")
        wait_js(page, "document.querySelector('#commentsLog').textContent.includes('댓글 수집: 1건')")
        page.locator("#workBar").wait_for(state="hidden")
        assert page.inner_text("#recollectSel") == "선택한 글 댓글 다시 수집 (0)"
        assert "댓글 7개" in page.inner_text("#commentsSummary")
        raw_xlsx = save_download("#dlComments")
        wb = load_workbook(raw_xlsx)
        assert [c.value for c in wb["언론"][2]] == ["domain", "text", "작성일", "게시글 URL"]  # 성향 분석 전 파일엔 성향 열 없음
        assert wb["언론"]["B3"].font.name == "현대산스 Text Medium"
        assert "A1:D1" in [str(r) for r in wb["언론"].merged_cells.ranges]
        assert not wb["언론"].auto_filter.ref and not list(wb["언론"].conditional_formatting)  # 필터·성향 색 없음
        assert wb["언론"]["A1"].value == "CN8_가격_공개_언론 댓글"
        assert wb["언론"].max_row == 5 and wb["언론"].max_column == 4
        assert wb["언론"]["C3"].alignment.horizontal == "center"  # 작성일 가운데 정렬
        assert wb["언론"]["D2"].border.right.style == "thin" and wb["언론"]["D3"].border.right.style == "thin"  # 표 오른쪽 끝 테두리 (다른 칸과 같은 두께)
        assert wb["언론"].column_dimensions["D"].width >= len("https://n.news.naver.com/mnews/article/001/0000000001")  # URL 길이에 맞춘 너비

        # 4. 성향 분석
        page.click("#goSentiment")
        page.click("#runSentiment")
        page.locator("#sentimentSummary", has_text="분석한 댓글 7개").wait_for(timeout=60000)
        senti_xlsx = save_download("#dlSentiment")
        wb = load_workbook(senti_xlsx)
        got = {r[1]: r[4] for ws in wb.worksheets for r in ws.iter_rows(min_row=3, values_only=True)}
        cfs = [str(r.sqref) for r in wb["언론"].conditional_formatting]
        assert cfs == ["E1:E1048576"]  # 성향 색도 성향 열에
        assert [c.value for c in wb["언론"][2]] == ["domain", "text", "작성일", "게시글 URL", "성향"]  # 성향은 맨 오른쪽
        assert wb["언론"]["E2"].font.name == "현대산스 Text Medium" and not wb["언론"].auto_filter.ref
        assert got["현기 배짱장사 답이 없네"] == "단순비난"
        assert got["가격이 너무 비싸요 4천은 심했다"] == "부정"
        assert got["디자인 좋네요 기대됩니다"] == "긍정"
        # 목적·이슈명·강조 규칙이 프롬프트에 들어감
        sys_senti = next(c["system"] for c in claude_calls if "댓글 성향 분석" in c["system"])
        assert "CN8 가격 공개 이후 가격·옵션 여론 파악" in sys_senti and "가장 중요한 규칙" in sys_senti
        sys_rel = next(c["system"] for c in claude_calls if "게시물 관련성 판독" in c["system"])
        assert "CN8 가격 공개 이후 가격·옵션 여론 파악" in sys_rel and "가장 중요한 규칙" in sys_rel

        # 4-1. 손으로 판정 고치기 → 교정 사례로 기록, 지시사항을 주고 재분석 → 명령 규칙 메모리에 저장
        row = page.locator("#cmtTable tbody tr", has_text="디자인 좋네요 기대됩니다")
        row.locator("select").select_option("중립")
        page.click("#reBox summary")
        page.fill("#reInstr", "가격 때문에 다른 차를 사겠다는 댓글은 부정으로 본다")
        n_calls = len(claude_calls)
        page.click("#reRun")
        page.locator("#reResult", has_text="재분석").wait_for(timeout=60000)
        assert "메모리에 저장" in page.inner_text("#reResult")
        sys_re = [c["system"] for c in claude_calls[n_calls:] if "댓글 성향 분석" in c["system"]]
        assert sys_re and "가격 때문에 다른 차를 사겠다는 댓글은 부정으로 본다" in sys_re[0] and "명령 규칙 메모리" in sys_re[0]
        assert "“디자인 좋네요 기대됩니다” → 중립" in sys_re[0]  # 교정 사례
        re_user = "".join(c["user"] for c in claude_calls[n_calls:] if "댓글 성향 분석" in c["system"])
        assert "디자인 좋네요 기대됩니다" not in re_user  # 손으로 고친 판정은 그대로 둠
        assert row.locator("select").input_value() == "중립"
        page.click('nav .tab[data-tab="prompt"]')
        assert "가격 때문에 다른 차를 사겠다는 댓글은 부정으로 본다" in page.input_value("#memoryMd")
        assert "가장 자극적인 댓글 위주" in page.input_value("#memoryMd")  # 댓글 샘플 규칙이 메모리에 기본으로
        assert "디자인 좋네요 기대됩니다" in page.inner_text("#corrList")
        md = save_download("#dlMemory")
        assert "## 성향 교정 사례 (자동 기록)" in md.read_text(encoding="utf-8")

        # 5. 보고서
        page.click('nav .tab[data-tab="sentiment"]')
        page.click("#goReport")
        if os.environ.get("SHOT_DIR"):
            page.screenshot(path=os.path.join(os.environ["SHOT_DIR"], "report-summary.png"))
        summary_box = page.inner_text("#reportSummary")
        assert "2026-08-05 07:00 ~ 2026-08-06 10:00" in summary_box
        assert "CN8 가격, 아반떼 가격" in summary_box and "|" not in page.inner_text("#reportSummary .rsum-kw") and "검색어 2개" in summary_box
        assert re.search(r"후보 \d+건 중 관련 게시글 5건 확인", summary_box) and "커뮤니티 1/2" in summary_box, [l for l in summary_box.splitlines() if "후보" in l or "수집 결과" in l]
        assert page.locator("#reportSummary .rsum-kw b").count() == 2
        assert "커뮤니티" in summary_box and "부정+단순비난" in summary_box
        page.click("#writeReport")
        page.locator("#reportMsg", has_text="초안").wait_for(timeout=60000)
        assert page.input_value("#rTitle") == "CN8 가격 공개 관련 여론 동향"
        # id 가 아닌 글로 돌려줘도 찾아 넣고, 빈 칸은 센 댓글 순으로 채운다
        q = lambda s, k: page.input_value(f'.qcol[data-s="{s}"] [data-k="{k}"]')
        assert q("부정", "short") and q("부정", "long") and q("긍정", "short") and q("단순비난", "short"), page.inner_text("#reportMsg")
        assert all(not re.match(r"c\d+", l) for l in q("부정", "short").splitlines())
        assert "긍정 칸은 센 댓글 순으로 채웠습니다" in page.inner_text("#reportMsg")
        sys_rep = next(c["system"] for c in claude_calls if "★ 보고서 작성 규칙 ★" in c["system"])
        assert "가장 자극적이고 반응이 센 댓글" in sys_rep and "원문 그대로" in sys_rep
        # 보고서다운 용어·근거 있는 내용·많이 나온 논점 위주 요약 (기본 규칙 + 마크다운 메모리)
        assert "호구론 → 국내 구매자 차별 인식" in sys_rep and "근거 없는 내용 금지" in sys_rep
        assert "계약 취소·인수 거부 인증 확인'처럼 수집 자료에서 실제로 확인되지 않는 내용은 쓰지 않는다" in sys_rep  # 메모리 규칙
        user_rep = next(c["user"] for c in claude_calls if "★ 보고서 작성 규칙 ★" in c["system"])
        assert "[성향별 자주 나온 말" in user_rep and re.search(r"부정·단순비난 댓글 \d+개: .+\(\d+\)", user_rep), user_rep
        # 요약만 다시 쓰기: 주요 댓글 칸은 그대로 두고 제목·요약만 (후보 댓글을 적게 보냄)
        page.fill('.qcol[data-s="중립"] [data-k="short"]', "직접 고친 중립 댓글")
        page.fill("#rTitle", "임시 제목")
        n_calls = len(claude_calls)
        page.click("#writeSummary")
        page.locator("#reportMsg", has_text="요약만 다시 썼습니다").wait_for(timeout=60000)
        assert page.input_value("#rTitle") == "CN8 가격 공개 관련 여론 동향"
        assert page.input_value('.qcol[data-s="중립"] [data-k="short"]') == "직접 고친 중립 댓글"
        assert len(claude_calls) == n_calls + 1 and "주요 댓글은 고르지 말고" in claude_calls[-1]["user"]
        # 1페이지 댓글은 2페이지에 겹치지 않는다: AI 초안에서도 빼고, 손으로 겹치게 넣어도 PPT 에서 뺀다
        for sname in ["긍정", "부정", "단순비난"]:
            assert not set(q(sname, "short").splitlines()) & set(q(sname, "long").splitlines()), sname
        sys_dup = "quotes_short(1페이지)에 고른 댓글은 quotes_long(2페이지)에 다시 넣지 않는다" in sys_rep
        assert sys_dup and "2페이지 주요 댓글에 다시 넣지 않는다" in sys_rep and "나머지 댓글 중 중요하다고 판단되는 댓글" in sys_rep
        # 내용 없는 자음·ㅋㅋ·욕 줄임말 댓글은 주요 댓글 후보에서 뺀다
        assert page.evaluate("[meaningfulQuote('ㅣㄱㅋㅋㅋㅋㅋㅋㅋㅋㅋㅋㅋㅋㅋㅋㅋㅋㅅㅂ'), meaningfulQuote('ㅋㅋㅋ 현기 배짱장사 답이 없네')]") == [False, True]
        assert "내용이 없는 댓글은 고르지 않는다" in sys_rep and "내용이 없는 댓글은 주요 댓글에서 뺀다" in sys_rep
        neg_short = q("부정", "short").splitlines()
        page.fill('.qcol[data-s="부정"] [data-k="long"]', "\n".join(neg_short + ["@홍길동 2페이지 전용 부정 댓글😂👍🏻"]))
        # 중립은 밋밋한 댓글로: AI 규칙·메모리 + AI 없이 다시 고르는 버튼
        assert "최대한 중립스러운 댓글" in sys_rep and "1페이지는 이번 이슈와 관계있는 성의 있는 내용" in sys_rep  # 기본 규칙 + 메모리
        assert page.evaluate("stripMentions('@user-abc12 @김철수 가격이 문제 a@b.com')") == "가격이 문제 a@b.com"
        assert "'@닉네임' 부분은 지운다" in sys_rep
        assert page.evaluate("[laughHeavy('맨날 풀옵으로 비교하냐 ㅋㅋㅋㅋㅋㅋㅋㅋㅋㅋ'), laughHeavy('가격 좀 비싸네요 ㅋㅋ')]") == [True, False]
        assert "'ㅋㅋㅋㅋㅋㅋ', 'ㅎㅎㅎㅎㅎ'가 많이 들어간 댓글은 주요 댓글에서 가능한 한 제외" in sys_rep
        assert "정말 별 의미 없는" not in sys_rep  # 예전 메모리 문장은 지워짐
        assert "깡통 → 기본 모델" in sys_rep and "'깡통' 대신 '기본 모델'" in sys_rep
        assert "summary 에는 중립 여론·중립 비율을 쓰지 않는다" in sys_rep and "요약에는 중립 여론·중립 비율을 언급하지 않는다" in sys_rep
        # 밋밋하되 이슈와 관계있고 성의 있는 댓글 > 한두 마디 > 자극적인 댓글
        scores = page.evaluate("(() => { const t = issueTerms(); return ['CN8 가격 4천 초반이면 투싼이랑 비슷한 수준이네요', '그렇군요', '현기 배짱장사 답이 없네 ㅋㅋ'].map((x) => neutralScore(x, t)); })()")
        assert scores[0] > scores[1] and scores[0] > scores[2], scores
        assert page.evaluate("[calmQuote('현기 배짱장사 답이 없네 ㅋㅋ'), calmQuote('다음 달 출고 예정이라고 하네요')]") == [False, True]
        n_calls = len(claude_calls)
        page.click('.refill[data-s="중립"]')
        page.locator("#reportMsg", has_text="중립 댓글을 다시 골랐습니다").wait_for()
        assert q("중립", "short") and len(claude_calls) == n_calls
        # 다른 성향도 성향 색 버튼으로 AI 없이 다시 고르기 (1·2페이지 겹치지 않음)
        assert page.evaluate("getComputedStyle(document.querySelector('.refill[data-s=\"부정\"]')).backgroundColor") == "rgb(217, 48, 37)"
        bg = page.evaluate("[...document.querySelectorAll('.qcol')].map((e) => getComputedStyle(e).backgroundColor)")
        assert len(set(bg)) == 4, bg  # 성향별로 흐린 배경색이 다름
        page.click('.refill[data-s="부정"]')
        page.locator("#reportMsg", has_text="부정 댓글을 다시 골랐습니다 — 센 댓글부터").wait_for()
        assert q("부정", "short") and not set(q("부정", "short").splitlines()) & set(q("부정", "long").splitlines())
        assert len(claude_calls) == n_calls
        # 다시 고른 뒤 손으로 1페이지 댓글을 2페이지에 겹치게 넣어도 PPT 에선 빠진다
        neg_short = q("부정", "short").splitlines()
        page.fill('.qcol[data-s="부정"] [data-k="long"]', "\n".join(neg_short + ["@홍길동 2페이지 전용 부정 댓글😂👍🏻"]))
        page.dispatch_event('.qcol[data-s="부정"] [data-k="long"]', "change")
        page.fill('.qcol[data-s="중립"] [data-k="short"]', "직접 고친 중립 댓글🤣")
        page.dispatch_event('.qcol[data-s="중립"] [data-k="short"]', "change")
        # 요약이 PPT 한 줄(60자)을 넘으면 화면에 경고, PPT 에선 글자를 줄여 넣는다
        long_sum = "테스트 요약 첫째 문장 확인 中 " + "가격 부담과 옵션 구성 불만이 이어지는 상황 " * 4
        page.fill("#rSummary", f"{long_sum}\n{long_sum}")
        page.dispatch_event("#rSummary", "input")
        page.dispatch_event("#rSummary", "change")
        assert "PPT 한 줄 넘음" in page.inner_text("#rSummaryLen")
        pptx = save_download("#dlPptx")
        # 최종 저장: 저장할 폴더를 묻는다 (취소하면 저장 안 함) → 고른 폴더에 모든 파일
        page.evaluate("() => { window.showDirectoryPicker = async () => { throw new DOMException('cancel', 'AbortError'); }; }")
        page.click("#saveFinal")
        page.locator("#finalMsg", has_text="폴더를 고르지 않아 취소").wait_for()
        page.evaluate("""() => { window.__saved = [];
          window.showDirectoryPicker = async (opt) => ({ name: '보고서폴더', getFileHandle: async (n, o) => ({ createWritable: async () => ({
            write: async (b) => window.__saved.push([n, (b.byteLength ?? b.length) > 1000, opt.mode]), close: async () => {} }) }) }); }""")
        page.click("#saveFinal")
        page.locator("#finalMsg", has_text="최종 DB 저장 완료").wait_for(timeout=60000)
        assert "'보고서폴더' 폴더에 저장" in page.inner_text("#finalMsg")
        saved = page.evaluate("window.__saved")
        assert len(saved) == 5 and all(big and mode == "readwrite" for _, big, mode in saved), saved
        assert any(n.endswith("_여론 동향 보고서.pptx") for n, _, _ in saved) and any(n.endswith("_성향분석 포함.xlsx") for n, _, _ in saved)

        # 보고서 DB · 백업
        page.click('nav .tab[data-tab="db"]')
        page.locator("#dbTable .badge.final").wait_for()
        backup = save_download("#dbExport")

        # 예전 버전으로 읽은 언론 글이 있으면 '언론만 다시 수집' 안내
        page.click('nav .tab[data-tab="posts"]')
        page.evaluate("job.posts.filter(p => p.external).forEach(p => { p.source = '새 창 열림'; }); renderPosts();")
        assert page.is_visible("#oldNewsNotice") and "언론만 체크" in page.inner_text("#oldNewsNotice")
        # 예전 버전으로 판독해 게시글 성향이 빈 글 → '채우기' 버튼 (체크는 그대로)
        page.click('nav .tab[data-tab="posts"]')
        blank_yt = "job.posts.filter(p => p.channel === 'youtube').forEach(p => { p.postSentiment = ''; p.postSentimentManual = false; }); renderPosts();"
        # 체크하지 않은(무관) 글은 성향이 비어도 채울 대상이 아님
        page.evaluate("job.posts.filter(p => !p.checked).forEach(p => { p.postSentiment = ''; }); renderPosts();")
        assert page.is_hidden("#fillPostSenti")
        page.evaluate(blank_yt)
        assert "체크한 글 중 성향이 비어 있는 1건" in page.inner_text("#fillPostSenti")
        assert "판정된" in page.inner_text("#postsSentiment")
        page.click("#fillPostSenti")
        wait_js(page, "job.posts.filter(p => p.channel === 'youtube').every(p => p.postSentiment)", timeout=30000)
        assert page.evaluate("job.posts.filter(p => p.channel === 'youtube').every(p => p.checked)")
        page.evaluate(blank_yt)

        # 언론만 다시 수집: 다른 채널의 게시글·댓글·성향 분석은 그대로 (성향이 빈 다른 채널 글은 성향만 채움)
        for c in ["cafe", "youtube", "community"]:
            box = page.locator(f'#runChannels input[value="{c}"]')
            if box.count():
                box.uncheck()
        page.click("#collectPosts")
        page.locator("#postsSummary", has_text="언론 새로 수집").wait_for(timeout=120000)
        # 다른 채널 글의 성향은 그대로 남아 있고, 빈 글이 없으니 '채우기' 버튼은 숨김
        assert page.locator("#postsTable tbody tr", has_text="아반떼 가격 진짜 비싸네요").locator("select").input_value() == "단순비난"
        assert page.locator("#fillPostSenti").is_hidden()
        assert page.is_hidden("#oldNewsNotice")  # 다시 수집하면 안내가 사라짐
        assert "성향이 빈 1건" in page.text_content("#postsLog")
        assert page.evaluate("job.posts.filter(p => p.channel === 'youtube').every(p => p.postSentiment)")
        page.click('nav .tab[data-tab="comments"]')
        assert "댓글 4개" in page.inner_text("#commentsSummary"), page.inner_text("#commentsSummary")
        page.click('nav .tab[data-tab="sentiment"]')
        assert "분석한 댓글 4개" in page.inner_text("#sentimentSummary")
        ctx.close()

    assert errors == [], errors
    with zipfile.ZipFile(pptx) as z:
        slide1 = z.read("ppt/slides/slide1.xml").decode()
        slide2 = z.read("ppt/slides/slide2.xml").decode()
        chart = z.read("ppt/charts/chart1.xml").decode()
    assert "CN8 가격 공개 관련 여론 동향" in slide1 and "테스트 요약 첫째 문장 확인 中" in slide1
    assert "현대하모니 L" in slide1  # 양식 서체 유지
    assert "현기 배짱장사 답이 없네" in slide1 and "현기 배짱장사 답이 없네" not in slide2  # 1페이지에만
    s2text = "".join(re.findall(r"<a:t>(.*?)</a:t>", slide2))
    assert "‘2페이지 전용 부정 댓글’" in slide2 and "홍길동" not in slide2, s2text  # @닉네임·이모지는 빼고
    assert "직접 고친 중립 댓글’" in slide1  # 이모지는 빼고 (까만 네모 방지)
    assert not re.search(r"[\U0001F000-\U0001FAFF\uFE0F]", slide1 + slide2)
    sum_shape = slide1[slide1.index('name="직사각형 31"'):]
    sum_shape = sum_shape[: sum_shape.index("</p:sp>")]
    assert 'sz="1400"' not in sum_shape and re.search(r'sz="1[0-3]\d\d"', sum_shape)  # 넘쳐서 글자 줄임
    # 그래프: 긍정 계열에도 수치 표시, 5% 미만 막대는 수치 숨김
    sers = re.findall(r"<c:ser>.*?</c:ser>", chart, re.S)
    assert len(sers) == 4 and all("<c:dLbls>" in x for x in sers)
    assert "<c:delete val=\"1\"/>" in chart
    for t in neg_short:
        assert t in slide1 and t not in slide2, t
    # 전체 7개: 긍정 2(좋네요·기대) / 부정 3 / 단순비난 2
    assert "28.6%" in slide1 and "42.9%" in slide1
    assert "<c:v>0.286</c:v>" in chart
    with zipfile.ZipFile(backup) as z:
        names = z.namelist()
        meta = json.loads(z.read("backup.json"))
    assert len(meta["jobs"]) == 1 and meta["jobs"][0]["final"]
    assert sorted(f["kind"] for f in meta["files"]) == ["comments", "posts", "postsSentiment", "pptx", "sentiment"]
    assert any(f["name"].endswith("_게시물 리스트_성향분석 포함.xlsx") for f in meta["files"])
    assert all(f["path"] in names for f in meta["files"])


# ---------------------------------------------------------------- 카페: 바로 열면 안 보이는 글 → 네이버 검색 경유
def cafe_full_article():
    return (
        '<html><head><title>아반떼 클럽 : 네이버 카페</title></head><body>'
        '<h3 class="title_text">아반떼 가격 너무 비싸요</h3><div class="article_info"><span class="date">2026.08.06. 09:00</span></div>'
        '<div class="se-main-container">가격이 너무 올랐네요</div>'
        '<div class="CommentBox"><ul>'
        '<li id="c1" class="CommentItem"><a class="comment_nickname">철수</a><span class="text_comment">진짜 비싸요 심했다</span><span class="comment_info_date">2026.08.06. 09:10</span></li>'
        '<li id="c2" class="CommentItem CommentItem--reply"><a class="comment_nickname">영희</a><span class="text_comment">현기 배짱장사</span><span class="comment_info_date">2026.08.06. 09:20</span></li>'
        "</ul></div></body></html>"
    )


def start_fake_naver(tmp_path: Path):
    import ssl
    import subprocess
    import threading
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

    cert, key = tmp_path / "cert.pem", tmp_path / "key.pem"
    subprocess.run(["openssl", "req", "-x509", "-newkey", "rsa:2048", "-nodes", "-days", "1", "-subj", "/CN=cafe.naver.com",
                    "-keyout", str(key), "-out", str(cert)], check=True, capture_output=True)
    requests = []

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            host = self.headers.get("Host", "")
            requests.append(host + self.path)
            if host.startswith("search.naver.com"):
                body = '<html><body><ul><li><a href="https://cafe.naver.com/avantemd/555?art=SEARCHKEY">아반떼 가격 너무 비싸요</a></li></ul></body></html>'
            elif self.path.startswith("/avantemd/555?art=SEARCHKEY"):
                body = cafe_full_article()
            elif self.path.startswith("/f-e/cafes/123/menus/0"):
                page = parse_qs(urlparse(self.path).query).get("page", ["1"])[0]
                body = (
                    '<html><head><title>아반떼 클럽 : 네이버 카페</title></head><body><table>'
                    '<tr><td><a class="article" href="https://cafe.naver.com/avantemd/555">아반떼 가격 너무 비싸요</a></td><td class="td_date">2026.08.06. 09:00</td></tr>'
                    "</table></body></html>"
                ) if page == "1" else "<html><body>검색 결과가 없습니다</body></html>"
            elif self.path.startswith("/f-e/cafes/123/articles/555") or self.path.startswith("/avantemd/555"):
                # 바로 열면 제목만 보이고 본문·댓글은 안 보임 (등급 제한)
                body = '<html><body><h3 class="title_text">아반떼 가격 너무 비싸요</h3><p>등급이 되어야 읽을 수 있습니다</p></body></html>'
            elif self.path.startswith("/avantemd"):
                body = '<html><script>var g_sClubId = "123";</script><body>카페 홈</body></html>'
            else:
                body = "<html><body></body></html>"
            data = body.encode()
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    tls = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
    tls.load_cert_chain(cert, key)
    server.socket = tls.wrap_socket(server.socket, server_side=True)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    server.requests = requests
    return server


def test_cafe_via_search(tmp_path: Path):
    server = start_fake_naver(tmp_path)
    port = server.server_address[1]
    claude_calls = []
    with sync_playwright() as p:
        ctx = p.chromium.launch_persistent_context(
            str(tmp_path / "profile"),
            executable_path=CHROMIUM,
            headless=True,
            timezone_id="Asia/Seoul",
            args=[
                f"--disable-extensions-except={EXT_DIR}",
                f"--load-extension={EXT_DIR}",
                f"--host-resolver-rules=MAP cafe.naver.com 127.0.0.1:{port}, MAP search.naver.com 127.0.0.1:{port}",
                "--no-proxy-server",
                "--ignore-certificate-errors",
            ],
            accept_downloads=True,
        )

        def anthropic(route):
            body = json.loads(route.request.post_data)
            out = json.dumps(claude(body, claude_calls), ensure_ascii=False)
            route.fulfill(status=200, headers={"content-type": "text/event-stream", "access-control-allow-origin": "*"}, body=sse(out))

        ctx.route("https://api.anthropic.com/**", anthropic)
        sw = ctx.service_workers[0] if ctx.service_workers else ctx.wait_for_event("serviceworker")
        ext_id = sw.url.split("/")[2]
        page = ctx.new_page()
        errors = []
        page.on("pageerror", lambda e: errors.append(str(e)))
        page.on("dialog", lambda d: d.accept())
        page.goto(f"chrome-extension://{ext_id}/app.html")
        page.click('nav .tab[data-tab="settings"]')
        page.fill("#apiKey", "sk-ant-test")
        page.click("#saveSettings")
        page.locator("#settingsMsg", has_text="저장").wait_for()
        page.click('nav .tab[data-tab="setup"]')
        page.fill("#purpose", "CN8 가격 공개 이후 가격·옵션 여론 파악")
        page.fill("#issueName", "CN8 가격 공개")
        page.fill("#rangeFrom", "2026-08-05T07:00")
        page.fill("#rangeTo", "2026-08-06T10:00")
        page.fill("#keywords", "아반떼 가격")
        for cb in ["#chNews", "#chYoutube", "#chCommunity"]:
            page.uncheck(cb)
        # 찾기 칸이 비어 있어도 동호회 목록이 보이고 '전체 선택'이 된다
        page.click("#cafeAll")
        assert page.locator("#cafeList input:checked").count() == page.locator("#cafeList input").count() > 40
        page.click("#cafeNone")
        page.fill("#cafeFilter", "아반떼 클럽")
        page.click("#cafeAll")
        page.fill("#cafeFilter", "")
        page.dispatch_event("#cafeFilter", "input")
        assert page.locator("#cafeList input:checked").count() == 1
        assert page.is_checked('#cafeList input[value="avantemd"]')
        page.click("#goCollect")
        page.click("#collectPosts")
        page.locator("#postsSummary", has_text="관련 게시글 1건").wait_for(timeout=180000)
        page.click("#goComments")
        page.click("#collectComments")
        page.locator("#commentsSummary", has_text="댓글을 모은 글 1/1").wait_for(timeout=180000)
        assert "댓글 2개" in page.inner_text("#commentsSummary")
        log = page.text_content("#commentsLog")
        ctx.close()
    server.shutdown()
    assert errors == [], errors
    assert "네이버 검색을 거쳐 읽음" in log, log
    assert any(r.startswith("search.naver.com") for r in server.requests), server.requests
    assert any("art=SEARCHKEY" in r for r in server.requests)


# ---------------------------------------------------------------- 뽐뿌: 댓글이 받아 온 화면에 없거나 구조를 몰라도 반드시 읽는다
PPOM_HEAD = '<html><head><meta charset="euc-kr"></head><body><div class="board-contents">아반떼 가격 너무 올랐네요 본문입니다</div>'


def ppom_view(no):
    if no == "980149":  # 댓글을 화면이 열린 뒤 불러옴 (받아 온 화면엔 없음)
        return PPOM_HEAD + (
            '<span>댓글 [3]</span><div id="cbox"></div><script>setTimeout(function(){var h="";'
            '["가격 실화냐 너무 비싸","옵션 장사 오지네","그래도 디자인은 좋음"].forEach(function(t,i){'
            'h+=\'<div class="comment_wrapper"><div id="commentContent_\'+i+\'" class="mid-text-area">\'+t+\'</div><span class="eng-day">2026-08-06 09:0\'+i+\'</span></div>\';});'
            'document.getElementById("cbox").innerHTML=h;},400);</script></body></html>'
        )
    if no == "980162":  # 처음엔 2개, '댓글 더보기'를 눌러야 나머지
        item = lambda t: f'<div class="comment_wrapper"><div class="mid-text-area">{t}</div></div>'
        return PPOM_HEAD + (
            '<span>댓글 [4]</span><div id="cbox">' + item("첫 댓글 비싸요") + item("둘째 댓글 기대돼요") + '</div>'
            '<a href="#" class="comment_more" onclick="more();return false">댓글 더보기</a>'
            "<script>function more(){document.getElementById('cbox').insertAdjacentHTML('beforeend',"
            + json.dumps(item("셋째 댓글 노조 탓") + item("넷째 댓글 그냥 그래"))
            + ");document.querySelector('.comment_more').remove();}</script></body></html>"
        )
    if no == "980135":  # 이름으로는 댓글인지 알 수 없는 구조 → AI 가 화면 글자에서
        return PPOM_HEAD + '<span>댓글 [2]</span><div class="zz"><p>닉네임A</p><p>댓글내용: 이 가격이면 안 산다</p></div><div class="zz"><p>닉네임B</p><p>댓글내용: 투싼이 낫겠네</p></div></body></html>'
    return PPOM_HEAD + '<span>댓글 [2]</span><div class="zz"><p>알 수 없는 구조</p></div></body></html>'


PPOM_LIST = (
    '<html><head><meta charset="euc-kr"></head><body><table>'
    '<tr class="baseList"><td><a href="view.php?id=car&amp;no=980149"><span class="baseList-title">아반떼 가격 실화냐</span></a>'
    '<span class="list_comment2"><span>18</span></span></td><td><time class="baseList-time">2026-08-06 08:00:00</time></td></tr>'
    '<tr class="baseList"><td><a href="view.php?id=car&amp;no=980100"><span>아반떼 가격 옛날 글</span></a></td><td>2026-07-01 08:00:00</td></tr>'
    "</table></body></html>"
)


def ppom_page(path):
    m = re.search(r"no=(\d+)", path)
    body = ppom_view(m.group(1)) if "view.php" in path and m else PPOM_LIST
    return "text/html; charset=euc-kr", body.encode("euc-kr")


def start_https(tmp_path: Path, host: str, page_fn):
    import ssl
    import subprocess
    import threading
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

    cert, key = tmp_path / "ppom-cert.pem", tmp_path / "ppom-key.pem"
    subprocess.run(["openssl", "req", "-x509", "-newkey", "rsa:2048", "-nodes", "-days", "1", "-subj", f"/CN={host}",
                    "-keyout", str(key), "-out", str(cert)], check=True, capture_output=True)

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            ctype, data = page_fn(self.path)
            self.send_response(200)
            self.send_header("Content-Type", ctype)
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    tls = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
    tls.load_cert_chain(cert, key)
    server.socket = tls.wrap_socket(server.socket, server_side=True)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    return server


def test_community_comments_must_collect(tmp_path: Path):
    claude_calls = []
    with sync_playwright() as p:
        # 탭으로 여는 화면에는 route 가 걸리지 않아, 진짜 주소를 내 컴퓨터의 가짜 서버로 돌린다
        server = start_https(tmp_path, "www.ppomppu.co.kr", ppom_page)
        ctx = p.chromium.launch_persistent_context(
            str(tmp_path / "profile"),
            executable_path=CHROMIUM,
            headless=True,
            timezone_id="Asia/Seoul",
            args=[
                f"--disable-extensions-except={EXT_DIR}",
                f"--load-extension={EXT_DIR}",
                f"--host-resolver-rules=MAP www.ppomppu.co.kr 127.0.0.1:{server.server_address[1]}",
                "--no-proxy-server",
                "--ignore-certificate-errors",
            ],
        )

        def anthropic(route):
            body = json.loads(route.request.post_data)
            out = json.dumps(claude(body, claude_calls), ensure_ascii=False)
            route.fulfill(status=200, headers={"content-type": "text/event-stream", "access-control-allow-origin": "*"}, body=sse(out))

        ctx.route("https://api.anthropic.com/**", anthropic)
        sw = ctx.service_workers[0] if ctx.service_workers else ctx.wait_for_event("serviceworker")
        ext_id = sw.url.split("/")[2]
        page = ctx.new_page()
        errors = []
        page.on("pageerror", lambda e: errors.append(str(e)))
        page.goto(f"chrome-extension://{ext_id}/app.html")
        wait_js(page, "typeof CommunityCollector === 'function' && typeof settings === 'object'")
        res = page.evaluate(
            """async () => {
              const logs = [];
              const make = (ai) => new CommunityCollector({ delays: { community: [0, 0] }, log: (m) => logs.push(m), isStopped: () => false,
                ai: ai ? new AiClient({ ...settings, apiKey: "sk-ant-test" }) : null });
              const mc = make(true);
              const site = { name: "뽐뿌", url: "https://www.ppomppu.co.kr/zboard/zboard.php?id=car" };
              const post = (no) => ({ siteUrl: site.url, url: `https://www.ppomppu.co.kr/zboard/view.php?id=car&no=${no}`, title: "아반떼 가격" });
              const out = {};
              const kw = parseKeywords("아반떼 가격");
              const found = await mc.search(site, kw, { from: new Date(2026, 7, 5), to: new Date(2026, 7, 7) });
              out.list = found.map((f) => [f.title, f.ccount]);
              for (const no of ["980149", "980162", "980135"]) {
                const a = await mc.readArticle(post(no));
                out[no] = { how: a.how, texts: a.comments.map((c) => c.text), dates: a.comments.map((c) => c.writtenAt) };
              }
              try { await make(false).readArticle(post("980170")); out.fail = "no error"; } catch (e) { out.fail = e.message; }
              out.logs = logs;
              return out;
            }"""
        )
        assert res["list"] == [["아반떼 가격 실화냐", 18]], res
        assert res["980149"]["how"] == "탭으로 열어 다시 읽음" and res["980149"]["texts"] == ["가격 실화냐 너무 비싸", "옵션 장사 오지네", "그래도 디자인은 좋음"], res
        assert res["980149"]["dates"][0] == "2026-08-06 09:00", res
        assert res["980162"]["texts"] == ["첫 댓글 비싸요", "둘째 댓글 기대돼요", "셋째 댓글 노조 탓", "넷째 댓글 그냥 그래"], res
        assert res["980135"]["how"] == "AI 화면 판독" and res["980135"]["texts"] == ["이 가격이면 안 산다", "투싼이 낫겠네"], res
        extract = [c for c in claude_calls if "★ 댓글 화면 판독 규칙 ★" in c["system"]]
        assert "한 글자도 바꾸거나" in extract[0]["system"] and "화면에 보이는 댓글 수: 2개" in extract[0]["user"]
        # 읽지 못하면 0개로 끝내지 않고 실패로 남겨 '못 읽은 글만 다시' 대상이 된다
        assert "댓글 2개가 보이는데 읽지 못했습니다" in res["fail"], res
        assert not errors, errors
        ctx.close()
        server.shutdown()
