"""네이버 카페 게시글 댓글 수집기.

사용법
  python naver_cafe_comments.py                 # 대화형 모드 (로그인 확인 → URL 입력 → 엑셀 저장)
  python naver_cafe_comments.py login           # 브라우저를 띄워 네이버에 직접 로그인 (최초 1회)
  python naver_cafe_comments.py collect URL ... # 지정한 게시글들의 댓글 수집
  python naver_cafe_comments.py collect -f urls.txt -o 결과.xlsx

로그인 세션은 ./naver_profile 폴더(브라우저 프로필)에 저장되어 다음 실행 때 재사용됩니다.
"""

from __future__ import annotations

import argparse
import re
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from urllib.parse import unquote

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from playwright.sync_api import BrowserContext, Page, sync_playwright

PROFILE_DIR = Path(__file__).resolve().parent / "naver_profile"
API_BASE = "https://apis.naver.com/cafe-web/cafe-articleapi/v2.1"
LOGIN_URL = "https://nid.naver.com/nidlogin.login?url=https%3A%2F%2Fcafe.naver.com%2F"
REQUEST_DELAY = 0.5  # 요청 사이 대기(초). 서버 부담을 줄이기 위해 너무 줄이지 마세요.
MAX_PAGES = 1000

# ----------------------------------------------------------------------------
# 데이터 구조
# ----------------------------------------------------------------------------


@dataclass
class Comment:
    comment_id: str
    is_reply: bool
    parent_id: str
    nickname: str
    writer_id: str
    content: str
    written_at: str
    is_deleted: bool = False
    attachment: str = ""


@dataclass
class ArticleResult:
    url: str
    cafe_id: str = ""
    article_id: str = ""
    title: str = ""
    author: str = ""
    comments: list[Comment] = field(default_factory=list)
    method: str = ""
    error: str = ""


# ----------------------------------------------------------------------------
# URL 해석
# ----------------------------------------------------------------------------

_CAFES_ARTICLES_RE = re.compile(r"/cafes/(\d+)/articles/(\d+)")
_CLUBID_RE = re.compile(r"clubid=(\d+)", re.I)
_ARTICLEID_RE = re.compile(r"articleid=(\d+)", re.I)
_SHORT_RE = re.compile(r"^https?://(?:m\.)?cafe\.naver\.com/([A-Za-z0-9_\-]+)/(\d+)(?:[/?#]|$)")


def parse_article_url(url: str) -> tuple[str | None, str | None]:
    """URL에서 (카페ID, 게시글ID)를 추출한다. 알 수 없는 값은 None.

    지원 형식
      https://cafe.naver.com/ca-fe/cafes/123/articles/456
      https://cafe.naver.com/f-e/cafes/123/articles/456
      https://m.cafe.naver.com/ca-fe/web/cafes/123/articles/456
      https://cafe.naver.com/ArticleRead.nhn?clubid=123&articleid=456
      https://cafe.naver.com/카페주소?iframe_url_utf8=...clubid%3D123%26articleid%3D456
      https://cafe.naver.com/카페주소/456          (카페ID는 페이지를 열어 확인)
    """
    url = url.strip()
    decoded = url
    for _ in range(3):  # iframe_url 은 여러 번 인코딩되어 있을 수 있음
        decoded = unquote(decoded)

    m = _CAFES_ARTICLES_RE.search(decoded)
    if m:
        return m.group(1), m.group(2)

    club = _CLUBID_RE.search(decoded)
    art = _ARTICLEID_RE.search(decoded)
    if club and art:
        return club.group(1), art.group(1)

    m = _SHORT_RE.match(url)
    if m:
        return None, m.group(2)

    return (club.group(1) if club else None), (art.group(1) if art else None)


def resolve_ids_via_browser(page: Page, url: str) -> tuple[str | None, str | None]:
    """게시글 페이지를 직접 열어서 네트워크 요청/iframe 주소로 카페ID·게시글ID를 알아낸다."""
    found: dict[str, str] = {}

    def on_request(req):
        if "cafe_id" in found:
            return
        m = _CAFES_ARTICLES_RE.search(req.url)
        if m:
            found["cafe_id"], found["article_id"] = m.group(1), m.group(2)

    page.on("request", on_request)
    try:
        page.goto(url, wait_until="domcontentloaded", timeout=30000)
        for _ in range(20):
            if "cafe_id" in found:
                break
            for fr in page.frames:
                cid, aid = parse_article_url(fr.url)
                if cid and aid:
                    found["cafe_id"], found["article_id"] = cid, aid
                    break
            if "cafe_id" in found:
                break
            page.wait_for_timeout(500)
    finally:
        page.remove_listener("request", on_request)
    return found.get("cafe_id"), found.get("article_id")


# ----------------------------------------------------------------------------
# API 응답 파싱
# ----------------------------------------------------------------------------


def _first(d: dict, *keys, default=""):
    for k in keys:
        v = d.get(k)
        if v not in (None, ""):
            return v
    return default


def format_timestamp(value) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, (int, float)) or (isinstance(value, str) and value.isdigit()):
        ts = float(value)
        if ts > 1e11:  # 밀리초
            ts /= 1000
        return datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M:%S")
    return str(value)


def parse_comment_item(item: dict) -> Comment:
    writer = item.get("writer") or {}
    cid = str(_first(item, "id", "commentId", "commentNo"))
    ref = str(_first(item, "refId", "refCommentId", "parentCommentId", default=""))
    is_reply = bool(item.get("isRef")) or bool(item.get("isReply")) or (ref != "" and ref != cid)

    attachments = []
    image = item.get("image") or {}
    if isinstance(image, dict) and image.get("url"):
        attachments.append(image["url"])
    sticker = item.get("sticker") or {}
    if isinstance(sticker, dict) and sticker.get("url"):
        attachments.append(sticker["url"])

    return Comment(
        comment_id=cid,
        is_reply=is_reply,
        parent_id=ref if is_reply else "",
        nickname=str(_first(writer, "nick", "nickname") or _first(item, "writerNickname", "nickname")),
        writer_id=str(_first(writer, "id", "memberId") or _first(item, "writerId", "memberId")),
        content=str(item.get("content") or ""),
        written_at=format_timestamp(_first(item, "updateDate", "writeDate", "createDate", default=None)),
        is_deleted=bool(item.get("isDeleted") or item.get("deleted")),
        attachment=" ".join(attachments),
    )


def extract_comment_items(payload: dict) -> list[dict] | None:
    """API 응답에서 댓글 목록을 찾는다. 형식을 알 수 없으면 None."""
    result = payload.get("result", payload)
    comments = result.get("comments")
    if isinstance(comments, dict):
        items = comments.get("items")
        if isinstance(items, list):
            return items
    if isinstance(comments, list):
        return comments
    if isinstance(result.get("items"), list):
        return result["items"]
    return None


# ----------------------------------------------------------------------------
# 수집 (API 방식)
# ----------------------------------------------------------------------------


class CollectError(Exception):
    pass


def _api_get(context: BrowserContext, url: str, referer: str) -> dict:
    resp = context.request.get(
        url,
        headers={
            "Referer": referer,
            "Accept": "application/json, text/plain, */*",
            "X-Cafe-Product": "pc",
        },
        timeout=30000,
    )
    if resp.status in (401, 403):
        raise CollectError(f"접근 권한 없음(HTTP {resp.status}) - 로그인 또는 카페 가입 여부를 확인하세요")
    if not resp.ok:
        raise CollectError(f"HTTP {resp.status}")
    try:
        return resp.json()
    except Exception as e:  # noqa: BLE001
        raise CollectError(f"JSON 응답이 아님: {e}") from e


def collect_via_api(context: BrowserContext, result: ArticleResult) -> None:
    cafe_id, article_id = result.cafe_id, result.article_id
    referer = f"https://cafe.naver.com/ca-fe/cafes/{cafe_id}/articles/{article_id}"

    # 게시글 정보(제목/작성자) - 실패해도 댓글 수집은 계속
    try:
        art = _api_get(
            context,
            f"{API_BASE}/cafes/{cafe_id}/articles/{article_id}?useCafeId=true&requestFrom=A",
            referer,
        )
        article = (art.get("result") or {}).get("article") or {}
        result.title = str(article.get("subject") or "")
        result.author = str((article.get("writer") or {}).get("nick") or "")
    except CollectError as e:
        if "권한" in str(e):
            raise

    seen: set[str] = set()
    comments: list[Comment] = []
    for page_no in range(1, MAX_PAGES + 1):
        data = _api_get(
            context,
            f"{API_BASE}/cafes/{cafe_id}/articles/{article_id}/comments/pages/{page_no}"
            f"?requestFrom=A&orderBy=asc",
            referer,
        )
        items = extract_comment_items(data)
        if items is None:
            raise CollectError("댓글 API 응답 형식을 인식하지 못함")
        new = 0
        for item in items:
            c = parse_comment_item(item)
            key = c.comment_id or f"{c.nickname}|{c.written_at}|{c.content}"
            if key in seen:
                continue
            seen.add(key)
            comments.append(c)
            new += 1
        if not items or new == 0:
            break
        time.sleep(REQUEST_DELAY)

    result.comments = comments
    result.method = "API"


# ----------------------------------------------------------------------------
# 수집 (화면 방식 - API가 실패했을 때 대체 수단)
# ----------------------------------------------------------------------------

_DOM_EXTRACT_JS = r"""
() => {
  const items = Array.from(document.querySelectorAll('li.CommentItem'));
  let lastParent = '';
  return items.map((li) => {
    const q = (sel) => { const el = li.querySelector(sel); return el ? el.innerText.trim() : ''; };
    const isReply = li.classList.contains('CommentItem--reply');
    const id = li.id || '';
    if (!isReply) lastParent = id;
    const img = li.querySelector('.comment_image img, .CommentItemImage img, .comment_sticker img, .se-sticker-image');
    return {
      id: id,
      isReply: isReply,
      parentId: isReply ? lastParent : '',
      nickname: q('.comment_nickname'),
      content: q('.text_comment'),
      date: q('.comment_info_date'),
      deleted: !!li.querySelector('.comment_deleted, .comment_delete'),
      attachment: img ? (img.getAttribute('src') || '') : '',
    };
  });
}
"""


def _find_comment_frame(page: Page):
    for fr in page.frames:
        try:
            if fr.query_selector("li.CommentItem, .CommentBox"):
                return fr
        except Exception:  # noqa: BLE001
            continue
    return None


def dom_rows_to_comments(rows: list[dict]) -> list[Comment]:
    return [
        Comment(
            comment_id=r.get("id", ""),
            is_reply=bool(r.get("isReply")),
            parent_id=r.get("parentId", ""),
            nickname=r.get("nickname", ""),
            writer_id="",
            content=r.get("content", ""),
            written_at=r.get("date", ""),
            is_deleted=bool(r.get("deleted")),
            attachment=r.get("attachment", ""),
        )
        for r in rows
    ]


def collect_via_dom(page: Page, result: ArticleResult) -> None:
    target = result.url
    if result.cafe_id and result.article_id:
        target = f"https://cafe.naver.com/ca-fe/cafes/{result.cafe_id}/articles/{result.article_id}"
    page.goto(target, wait_until="domcontentloaded", timeout=30000)

    frame = None
    for _ in range(30):
        frame = _find_comment_frame(page)
        if frame:
            break
        page.wait_for_timeout(500)
    if not frame:
        raise CollectError("화면에서 댓글 영역을 찾지 못함 (권한이 없거나 댓글이 없는 글일 수 있음)")

    if not result.title:
        el = frame.query_selector(".title_text, h3.title_text")
        if el:
            result.title = el.inner_text().strip()

    seen: set[str] = set()
    comments: list[Comment] = []
    visited_pages: set[str] = set()
    while True:
        frame.wait_for_timeout(800)
        for c in dom_rows_to_comments(frame.evaluate(_DOM_EXTRACT_JS)):
            key = c.comment_id or f"{c.nickname}|{c.written_at}|{c.content}"
            if key not in seen:
                seen.add(key)
                comments.append(c)

        # 댓글 페이지 이동 버튼(숫자) 중 아직 방문하지 않은 것을 클릭
        current = frame.query_selector(".CommentBox .ArticlePaginate button[aria-pressed='true']")
        if current:
            visited_pages.add(current.inner_text().strip())
        next_btn = None
        for btn in frame.query_selector_all(".CommentBox .ArticlePaginate button"):
            label = btn.inner_text().strip()
            if label.isdigit() and label not in visited_pages:
                next_btn = btn
                visited_pages.add(label)
                break
        if not next_btn:
            break
        next_btn.click()

    result.comments = comments
    result.method = "화면"


# ----------------------------------------------------------------------------
# 전체 흐름
# ----------------------------------------------------------------------------


def open_context(p, headless: bool) -> BrowserContext:
    PROFILE_DIR.mkdir(exist_ok=True)
    return p.chromium.launch_persistent_context(
        str(PROFILE_DIR),
        headless=headless,
        locale="ko-KR",
        viewport={"width": 1280, "height": 900},
    )


def is_logged_in(context: BrowserContext) -> bool:
    names = {c["name"] for c in context.cookies("https://naver.com")}
    return "NID_AUT" in names and "NID_SES" in names


def do_login(p) -> None:
    context = open_context(p, headless=False)
    page = context.pages[0] if context.pages else context.new_page()
    page.goto(LOGIN_URL)
    print("\n브라우저 창에서 네이버 로그인을 진행하세요.")
    print("('로그인 상태 유지'를 체크하면 다음부터 다시 로그인하지 않아도 됩니다)")
    for _ in range(600):  # 최대 5분 대기
        if is_logged_in(context):
            break
        time.sleep(0.5)
    if is_logged_in(context):
        print("로그인 완료! 세션이 저장되었습니다.")
    else:
        print("로그인이 확인되지 않았습니다. 다시 시도해 주세요.")
    context.close()


def collect(p, urls: list[str], headless: bool) -> list[ArticleResult]:
    context = open_context(p, headless=headless)
    if not is_logged_in(context):
        print("[경고] 로그인 세션이 없습니다. 비공개/멤버 전용 글은 수집되지 않을 수 있습니다.")
        print("       먼저 `python naver_cafe_comments.py login` 을 실행하세요.")
    page = context.pages[0] if context.pages else context.new_page()

    results: list[ArticleResult] = []
    for idx, url in enumerate(urls, 1):
        print(f"[{idx}/{len(urls)}] {url}")
        r = ArticleResult(url=url)
        try:
            cafe_id, article_id = parse_article_url(url)
            if not cafe_id:
                cafe_id, article_id = resolve_ids_via_browser(page, url)
            if not (cafe_id and article_id):
                raise CollectError("URL에서 카페ID/게시글번호를 찾지 못함")
            r.cafe_id, r.article_id = cafe_id, article_id

            try:
                collect_via_api(context, r)
            except CollectError as api_err:
                print(f"    API 수집 실패({api_err}) → 화면에서 수집 시도")
                collect_via_dom(page, r)
            print(f"    완료: 댓글 {len(r.comments)}개 ({r.method}) {r.title}")
        except Exception as e:  # noqa: BLE001
            r.error = str(e)
            print(f"    실패: {e}")
        results.append(r)
        time.sleep(REQUEST_DELAY)

    context.close()
    return results


# ----------------------------------------------------------------------------
# 엑셀 저장
# ----------------------------------------------------------------------------

COMMENT_HEADERS = [
    "게시글URL", "카페ID", "게시글번호", "게시글제목", "게시글작성자",
    "순번", "구분", "댓글ID", "부모댓글ID", "작성자닉네임", "작성자ID",
    "댓글내용", "작성일시", "삭제여부", "첨부(이미지/스티커)",
]
SUMMARY_HEADERS = ["게시글URL", "카페ID", "게시글번호", "게시글제목", "게시글작성자", "댓글수", "수집방식", "상태"]


def _clean(value):
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _style_sheet(ws, widths: dict[int, int]) -> None:
    header_fill = PatternFill("solid", fgColor="03C75A")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in widths.items():
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width


def save_excel(results: list[ArticleResult], path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "댓글"
    ws.append(COMMENT_HEADERS)
    for r in results:
        for i, c in enumerate(r.comments, 1):
            ws.append([_clean(v) for v in [
                r.url, r.cafe_id, r.article_id, r.title, r.author,
                i, "답글" if c.is_reply else "댓글", c.comment_id, c.parent_id,
                c.nickname, c.writer_id, c.content, c.written_at,
                "Y" if c.is_deleted else "", c.attachment,
            ]])
    for row in ws.iter_rows(min_row=2, min_col=12, max_col=12):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _style_sheet(ws, {1: 45, 4: 30, 10: 16, 12: 60, 13: 20, 15: 30})

    ws2 = wb.create_sheet("요약")
    ws2.append(SUMMARY_HEADERS)
    for r in results:
        ws2.append([_clean(v) for v in [
            r.url, r.cafe_id, r.article_id, r.title, r.author,
            len(r.comments), r.method, r.error or "성공",
        ]])
    _style_sheet(ws2, {1: 45, 4: 40, 8: 40})

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ----------------------------------------------------------------------------
# 명령행
# ----------------------------------------------------------------------------


def read_urls(args_urls: list[str], url_file: str | None) -> list[str]:
    urls = list(args_urls or [])
    if url_file:
        for line in Path(url_file).read_text(encoding="utf-8-sig").splitlines():
            line = line.strip()
            if line and not line.startswith("#"):
                urls.append(line)
    # 중복 제거(순서 유지)
    return list(dict.fromkeys(u.strip() for u in urls if u.strip()))


def default_output() -> Path:
    return Path(f"네이버카페_댓글_{datetime.now():%Y%m%d_%H%M%S}.xlsx")


def interactive(p) -> None:
    context = open_context(p, headless=True)
    logged_in = is_logged_in(context)
    context.close()
    if not logged_in:
        print("저장된 로그인 세션이 없습니다. 로그인 창을 엽니다.")
        do_login(p)

    print("\n댓글을 수집할 게시글 URL을 한 줄에 하나씩 입력하세요. (입력을 마치면 빈 줄에서 Enter)")
    urls = []
    while True:
        try:
            line = input("URL> ").strip()
        except EOFError:
            break
        if not line:
            break
        urls.append(line)
    urls = read_urls(urls, None)
    if not urls:
        print("입력된 URL이 없습니다.")
        return
    results = collect(p, urls, headless=True)
    out = save_excel(results, default_output())
    print(f"\n엑셀 저장 완료: {out.resolve()}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="네이버 카페 게시글 댓글 수집기 (엑셀 저장)")
    sub = parser.add_subparsers(dest="command")
    sub.add_parser("login", help="브라우저를 띄워 네이버에 로그인하고 세션 저장")
    c = sub.add_parser("collect", help="게시글 URL들의 댓글을 수집해 엑셀로 저장")
    c.add_argument("urls", nargs="*", help="게시글 URL (여러 개 가능)")
    c.add_argument("-f", "--file", help="URL 목록 텍스트 파일 (한 줄에 하나)")
    c.add_argument("-o", "--output", help="저장할 엑셀 파일 경로 (기본: 네이버카페_댓글_날짜시간.xlsx)")
    c.add_argument("--show-browser", action="store_true", help="수집 중 브라우저 창 표시")
    args = parser.parse_args(argv)

    with sync_playwright() as p:
        if args.command == "login":
            do_login(p)
        elif args.command == "collect":
            urls = read_urls(args.urls, args.file)
            if not urls:
                parser.error("URL을 입력하거나 -f 로 URL 파일을 지정하세요.")
            results = collect(p, urls, headless=not args.show_browser)
            out = save_excel(results, Path(args.output) if args.output else default_output())
            total = sum(len(r.comments) for r in results)
            failed = sum(1 for r in results if r.error)
            print(f"\n총 {len(results)}개 글, 댓글 {total}개 수집 (실패 {failed}건)")
            print(f"엑셀 저장 완료: {out.resolve()}")
        else:
            interactive(p)
    return 0


if __name__ == "__main__":
    sys.exit(main())
